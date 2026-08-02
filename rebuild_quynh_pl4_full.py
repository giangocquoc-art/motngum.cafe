from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BASE = Path(r"C:\Users\giang\OneDrive\Personal Vault\nghiên cứu khoa học\tiểu luận Quỳnh Phương x3\quynh-phuong-chu-de-4")
SOURCE = BASE / "output" / "K51NC_QuynhPhuong_ChuDe4_PL4.docx"
OUTPUT = BASE / "output" / "K51NC_QuynhPhuong_ChuDe4_hoan-chinh.docx"
ILLUSTRATION = BASE / "figures" / "auto" / "lop-hoc-hoa-nhap-minh-hoa-v2.png"


def set_run_font(run, size=12, bold=None, italic=None, all_caps=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.all_caps = all_caps


def set_body_format(p, first_line=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=0, line=1.5):
    pf = p.paragraph_format
    pf.alignment = align
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.left_indent = Cm(0)
    pf.right_indent = Cm(0)
    pf.first_line_indent = Cm(1.0) if first_line else Cm(0)
    for run in p.runs:
        set_run_font(run)
    return p


def add_body(doc, text, bold=False, italic=False, first_line=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
             before=0, after=0, keep=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run_font(r, 12, bold=bold, italic=italic)
    set_body_format(p, first_line=first_line, align=align, before=before, after=after)
    p.paragraph_format.keep_with_next = keep
    return p


def add_heading(doc, text, level=1, page_break=False):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.clear()
    r = p.add_run(text)
    if level == 1:
        set_run_font(r, 13, bold=True, all_caps=True)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(8)
    elif level == 2:
        set_run_font(r, 12, bold=True)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(3)
    else:
        set_run_font(r, 12, bold=True, italic=True)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = page_break
    return p


def add_front_title(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run_font(r, 14, bold=True, all_caps=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.keep_with_next = True
    return p


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    return p


def add_field_paragraph(doc, instruction, placeholder="Nhấn F9 để cập nhật trường tự động."):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.15
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r1 = p.add_run(); r1._r.append(begin)
    r2 = p.add_run(); r2._r.append(instr)
    r3 = p.add_run(); r3._r.append(separate)
    r4 = p.add_run(placeholder); set_run_font(r4, 12)
    r5 = p.add_run(); r5._r.append(end)
    return p


def add_seq_caption(doc, label, title, above=True):
    p = doc.add_paragraph(style="Caption" if "Caption" in [s.name for s in doc.styles] else None)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_before = Pt(4 if above else 2)
    p.paragraph_format.space_after = Pt(3 if above else 4)
    p.paragraph_format.keep_with_next = above
    r = p.add_run(f"{label} ")
    set_run_font(r, 12, bold=True)
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = f" SEQ {label} \\* ARABIC "
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    f1 = p.add_run(); f1._r.append(begin)
    f2 = p.add_run(); f2._r.append(instr)
    f3 = p.add_run(); f3._r.append(separate)
    f4 = p.add_run("1"); set_run_font(f4, 12, bold=True)
    f5 = p.add_run(); f5._r.append(end)
    tail = p.add_run(f". {title}")
    set_run_font(tail, 12, bold=True)
    return p


def shade(cell, fill="E7E6E6"):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def fill_cell(cell, text, bold=False, center=False, size=12):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.15
    pf.left_indent = Cm(0)
    pf.right_indent = Cm(0)
    pf.first_line_indent = Cm(0)
    r = p.add_run(text)
    set_run_font(r, size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    trPr.append(marker)


def prevent_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def set_table_geometry(table, widths_cm):
    widths = [int(Cm(w).emu / 635) for w in widths_cm]
    tblPr = table._tbl.tblPr
    tblW = tblPr.first_child_found_in("w:tblW")
    if tblW is None:
        tblW = OxmlElement("w:tblW"); tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(widths)))
    tblW.set(qn("w:type"), "dxa")
    tblInd = tblPr.first_child_found_in("w:tblInd")
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd"); tblPr.append(tblInd)
    tblInd.set(qn("w:w"), "0"); tblInd.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol"); col.set(qn("w:w"), str(width)); grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tcW = cell._tc.get_or_add_tcPr().tcW
            tcW.set(qn("w:w"), str(widths[i])); tcW.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths_cm, center_cols=(), font_size=12):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    for i, label in enumerate(headers):
        shade(header.cells[i])
        fill_cell(header.cells[i], label, bold=True, center=True, size=font_size)
    repeat_header(header)
    prevent_split(header)
    for values in rows:
        row = table.add_row()
        prevent_split(row)
        for i, value in enumerate(values):
            fill_cell(row.cells[i], str(value), center=i in center_cols, size=font_size)
    set_table_geometry(table, widths_cm)
    return table


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        set_run_font(r, 12)
        pf = p.paragraph_format
        pf.line_spacing = 1.5
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.left_indent = Cm(0.9)
        pf.first_line_indent = Cm(-0.45)


doc = Document(SOURCE)
body = doc._element.body

# Preserve the cover exactly, including its explicit page break, and retain the
# original section break that starts logical page numbering at the report.
section_break_para = next(
    deepcopy(p._p) for p in doc.paragraphs
    if p._p.pPr is not None and p._p.pPr.sectPr is not None
)
cover_break = next(p for p in doc.paragraphs if p._p.xpath('.//w:br[@w:type="page"]'))
cover_index = list(body).index(cover_break._p)
for element in list(body)[cover_index + 1:]:
    if element.tag != qn("w:sectPr"):
        body.remove(element)

# FRONT MATTER
add_front_title(doc, "LỜI CAM ĐOAN")
add_body(doc, (
    "Tôi cam đoan bài tiểu luận này được thực hiện trên cơ sở nghiên cứu tài liệu, phân tích yêu cầu của học phần "
    "và thiết kế sản phẩm giáo dục. Các nhận định, bảng biểu và tình huống minh họa được trình bày nhằm phục vụ "
    "mục đích học thuật; những nguồn tham khảo được ghi rõ ở cuối bài."
))
add_body(doc, (
    "Sản phẩm chưa thay thế kết quả khảo sát tại một cơ sở giáo dục cụ thể. Tình huống minh họa không mô tả một "
    "học sinh có thật và không được dùng để chẩn đoán, gắn nhãn hoặc xếp hạng giáo viên."
))
add_body(doc, "Học viên", first_line=False, align=WD_ALIGN_PARAGRAPH.RIGHT, before=8)
add_front_title(doc, "LỜI CẢM ƠN")
add_body(doc, (
    "Tôi trân trọng cảm ơn TS. Nguyễn Thị Quý đã định hướng việc lựa chọn vấn đề giáo dục hòa nhập và yêu cầu "
    "chuyển kiến thức của học phần thành một sản phẩm có thể sử dụng trong thực tiễn giáo dục tiểu học."
))
add_body(doc, (
    "Tôi cũng cảm ơn các đồng nghiệp đã chia sẻ những khó khăn thường gặp khi thiết kế nhiệm vụ, tổ chức sự tham "
    "gia của học sinh và phối hợp với gia đình. Những trao đổi đó giúp khung năng lực và các minh họa trong bài "
    "được diễn đạt theo hướng cụ thể, khả thi và tôn trọng sự đa dạng của người học."
))
add_page_break(doc)

add_front_title(doc, "MỤC LỤC")
add_field_paragraph(doc, ' TOC \\o "1-2" \\h \\z \\u ')
add_page_break(doc)

add_front_title(doc, "DANH MỤC BẢNG")
add_field_paragraph(doc, ' TOC \\h \\z \\c "Bảng" ')
add_front_title(doc, "DANH MỤC HÌNH")
add_field_paragraph(doc, ' TOC \\h \\z \\c "Hình" ')
add_front_title(doc, "DANH MỤC TỪ VIẾT TẮT")
add_table(doc, ["Từ viết tắt", "Nội dung"], [
    ("GDHN", "Giáo dục hòa nhập"),
    ("GV", "Giáo viên"),
    ("HS", "Học sinh"),
    ("UDL", "Thiết kế phổ quát cho học tập (Universal Design for Learning)"),
], [4.0, 12.5], center_cols=(0,))

# Insert the retained section break so the report begins on logical page 1.
body.sectPr.addprevious(section_break_para)

# REPORT
add_heading(doc, "PHẦN 1. GIỚI THIỆU VẤN ĐỀ", 1)
add_heading(doc, "1.1. Lý do lựa chọn chủ đề", 2)
add_body(doc, (
    "Giáo dục hòa nhập không chỉ là bố trí học sinh có nhu cầu hỗ trợ vào cùng một lớp học. Bản chất của cách tiếp "
    "cận này là nhận diện và tháo gỡ những rào cản khiến một số học sinh khó tiếp cận nội dung, khó tham gia hoạt "
    "động hoặc khó thể hiện kết quả học tập. Trong lớp tiểu học, rào cản có thể xuất hiện từ ngôn ngữ hướng dẫn, "
    "học liệu, cách tổ chức nhóm, không gian, công nghệ, nhịp độ bài học hoặc thái độ của người lớn và bạn học."
))
add_body(doc, (
    "Trong bối cảnh quốc tế hóa giáo dục, quyền được học tập, sự công bằng, tôn trọng khác biệt và năng lực hợp tác "
    "đa bên ngày càng trở thành những yêu cầu cốt lõi. UNESCO (2017) nhấn mạnh việc chuyển trọng tâm từ việc xem "
    "người học là vấn đề sang xem hệ thống và môi trường học tập là nơi cần được điều chỉnh. Sự chuyển dịch đó làm "
    "thay đổi vai trò của giáo viên: từ truyền đạt theo một cách thống nhất sang thiết kế nhiều con đường tham gia, "
    "theo dõi tiến bộ và điều chỉnh quyết định từ minh chứng."
))
add_body(doc, (
    "Vì vậy, quốc tế hóa trong chủ đề này không đồng nghĩa với đưa thật nhiều thuật ngữ hoặc thiết bị từ bên ngoài "
    "vào lớp học. "
    "Một kỹ thuật chỉ có ý nghĩa khi giải thích được rào cản cần xử lý và điều kiện để duy trì. Đây là lý do bài viết "
    "chọn khung năng lực thay vì tập hợp mẹo dạy học rời rạc."
))
add_heading(doc, "1.2. Mục đích, đối tượng và phạm vi", 2)
add_body(doc, (
    "Bài tiểu luận phân tích yêu cầu năng lực giáo dục hòa nhập của giáo viên phổ thông, tập trung vào khả năng vận "
    "dụng ở tiểu học. Sản phẩm trung tâm là một khung gồm năm năng lực, các chỉ số mô tả, biện pháp phát triển và "
    "một bộ minh họa sử dụng khung trong tình huống lớp học. Khung phục vụ tự rà soát và phát triển chuyên môn, "
    "không phải thang điểm kiểm định hay công cụ chẩn đoán học sinh."
))
add_heading(doc, "1.3. Phương pháp thực hiện", 2)
add_body(doc, (
    "Bài viết sử dụng phương pháp nghiên cứu tài liệu, phân tích quy định và đối chiếu các khung năng lực có liên "
    "quan. Kết quả đối chiếu được chuyển hóa thành chỉ số hành vi có thể quan sát. Tình huống minh họa "
    "được xây dựng giả định để kiểm tra tính rõ ràng, tính khả thi và sự nhất quán giữa nhận diện rào cản, lựa chọn "
    "hỗ trợ, thu thập minh chứng và điều chỉnh dạy học."
))

add_heading(doc, "PHẦN 2. CƠ SỞ LÝ LUẬN", 1)
add_heading(doc, "2.1. Các khái niệm liên quan", 2)
add_body(doc, (
    "Giáo dục hòa nhập là quá trình bảo đảm mọi người học được hiện diện, tham gia, tiến bộ và được tôn trọng trong "
    "một môi trường giáo dục chung. Đa dạng người học bao gồm khác biệt về khả năng, ngôn ngữ, văn hóa, hoàn cảnh, "
    "sức khỏe, nhịp độ và cách học. Rào cản học tập và tham gia là những yếu tố của nhiệm vụ, học liệu, giao tiếp, "
    "môi trường hoặc quan hệ xã hội làm giảm cơ hội học tập của một cá nhân hay một nhóm."
))
add_body(doc, (
    "Năng lực giáo dục hòa nhập của giáo viên là sự huy động có trách nhiệm kiến thức, kỹ năng, thái độ và phán đoán "
    "nghề nghiệp. Sự huy động đó thể hiện qua việc nhận diện rào cản, thiết kế cơ hội học tập tiếp cận được, tổ chức hỗ trợ, "
    "đánh giá vì sự tiến bộ và phối hợp với các bên liên quan. Năng lực không được suy ra từ một hành vi đơn lẻ mà cần được "
    "đọc qua chuỗi minh chứng trong bối cảnh cụ thể."
))
add_body(doc, (
    "Cần phân biệt bình đẳng, công bằng và hòa nhập. Bình đẳng thường cung cấp cùng một điều kiện cho tất cả học "
    "sinh. Công bằng cho phép phân bổ hỗ trợ khác nhau để các em có cơ hội đạt mục tiêu. Hòa nhập đòi hỏi thiết "
    "kế môi trường chung sao cho sự khác biệt không trở thành lý do bị tách khỏi hoạt động. Chẳng hạn, cùng yêu cầu "
    "xác định ý chính, học sinh có thể đọc bản chữ rõ, nghe giáo viên đọc hoặc sử dụng thẻ trình tự. Tiêu chí "
    "đánh giá vẫn tập trung vào việc hiểu ý chính chứ không chấm ưu thế của một hình thức thể hiện."
))
add_heading(doc, "2.2. Giáo dục hòa nhập trong bối cảnh quốc tế hóa", 2)
add_body(doc, (
    "Quốc tế hóa tạo điều kiện tiếp cận các nguyên tắc dựa trên quyền, thiết kế phổ quát cho học tập, đánh giá hình "
    "thành và công nghệ hỗ trợ. Đồng thời, việc tiếp nhận kinh nghiệm quốc tế phải gắn với chương trình, nguồn lực, "
    "văn hóa nhà trường và quy định Việt Nam. Sao chép một mô hình mà không xét điều kiện triển khai có thể làm tăng "
    "tải cho giáo viên hoặc tạo thêm bất bình đẳng giữa học sinh có và không có thiết bị."
))
add_body(doc, (
    "Kaplan và Bista (2022) gợi ý giáo viên quan sát sự tham gia thực tế của người học và điều chỉnh môi trường trước "
    "khi quy nguyên khó khăn cho cá nhân. Quan điểm này tương thích với quy định về giáo dục hòa nhập ở Việt Nam: "
    "tôn trọng đặc điểm người học, phối hợp với gia đình và các lực lượng hỗ trợ, đồng thời duy trì mục tiêu giáo dục "
    "phù hợp."
))
add_body(doc, (
    "Đối chiếu kinh nghiệm quốc tế với bối cảnh Việt Nam tạo ra hai yêu cầu. Giáo viên cần biết các chuẩn mực chung "
    "về quyền tham gia, điều chỉnh hợp lý và bảo mật; đồng thời phải xét sĩ số, học liệu, thời gian, đầu mối hỗ trợ và "
    "khả năng phối hợp của gia đình. Nhờ vậy, khung vẫn giữ định hướng quốc tế nhưng không trở thành một bộ yêu cầu "
    "xa rời công việc hằng ngày."
))
add_heading(doc, "2.3. Nguyên tắc thiết kế khung năng lực", 2)
add_body(doc, (
    "Khung được xây dựng theo bốn nguyên tắc. Thứ nhất, lấy quyền, phẩm giá và tiếng nói của học sinh làm điểm xuất "
    "phát. Thứ hai, chỉ số phải mô tả hành động nghề nghiệp có thể quan sát và tạo được minh chứng. Thứ ba, năng lực "
    "phải bao quát chu trình trước, trong và sau bài học thay vì chỉ nhấn mạnh kỹ thuật đứng lớp. Thứ tư, công nghệ, "
    "trong đó có AI, chỉ là phương tiện; giáo viên phải kiểm chứng đầu ra, bảo vệ dữ liệu và luôn có phương án không "
    "phụ thuộc thiết bị."
))
add_body(doc, (
    "Một chỉ số chỉ được giữ lại khi trả lời được ba câu hỏi: giáo viên làm gì, dựa vào thông tin nào và hành động đó "
    "thay đổi cơ hội học tập ra sao. Cách viết này hạn chế các mô tả khó kiểm chứng như 'có tinh thần hòa nhập'. Nó cũng tạo "
    "đường nối tới minh chứng. Kế hoạch cho thấy ý định thiết kế, quan sát cho thấy mức tham gia, sản phẩm trước-sau "
    "cho thấy tiến bộ, còn phản hồi học sinh kiểm tra tính hữu ích của hỗ trợ."
))

add_heading(doc, "PHẦN 3. THỰC TIỄN VÀ NHU CẦU", 1)
add_heading(doc, "3.1. Bối cảnh thực tiễn", 2)
add_body(doc, (
    "Trong lớp tiểu học, giáo viên thường đồng thời xử lý chênh lệch về vốn từ, tốc độ đọc, khả năng tập trung, vận "
    "động, giao tiếp và mức hỗ trợ từ gia đình. Một nhiệm vụ tưởng như đơn giản có thể đặt ra nhiều rào cản: văn bản "
    "dài nhưng không có gợi ý hình ảnh; yêu cầu trả lời chỉ bằng viết; thời gian hoàn thành cố định; hoặc cách chia "
    "nhóm khiến một học sinh luôn đứng ngoài hoạt động."
))
add_body(doc, (
    "Bài viết không sử dụng số liệu khảo sát của một đơn vị nên không đưa ra tỷ lệ hay kết luận thực trạng giả định. "
    "Từ những yêu cầu nghề nghiệp phổ biến, có thể xác định các nội dung cần được rà soát tại trường. Đó là mức hiểu về "
    "giáo dục hòa nhập, khả năng nhận diện rào cản, chất lượng điều chỉnh nhiệm vụ, cách theo dõi tiến bộ, cơ chế "
    "phối hợp với gia đình và việc bảo vệ dữ liệu học sinh."
))
add_body(doc, (
    "Nếu triển khai tại trường, khảo sát điểm xuất phát nên dùng một mẫu nhỏ nhưng có nhiều nguồn: kế hoạch bài dạy, "
    "học liệu, sản phẩm học sinh đã ẩn danh, quan sát hoạt động và trao đổi với giáo viên. Các nguồn phải được đọc cùng "
    "nhau. Kế hoạch ghi 'phân hóa' chưa chứng minh lớp học hòa nhập nếu một số em vẫn không có vai trò; ngược lại, một "
    "tiết học chưa trôi chảy cũng chưa đủ để kết luận khi điều kiện hỗ trợ chưa rõ."
))
add_heading(doc, "3.2. Cơ hội và thách thức", 2)
add_body(doc, (
    "Chương trình giáo dục phổ thông theo định hướng phát triển phẩm chất, năng lực tạo không gian cho giáo viên sử "
    "dụng nhiều hình thức tổ chức và đánh giá. Sinh hoạt tổ chuyên môn, dự giờ theo hướng nghiên cứu bài học và kho "
    "học liệu số cũng là điều kiện thuận lợi để chia sẻ điều chỉnh. Tuy nhiên, giáo viên có thể gặp hạn chế về thời "
    "gian, học liệu tiếp cận được, nhân lực hỗ trợ và sự thống nhất trong phối hợp."
))
add_body(doc, (
    "Một thách thức khác là nhầm lẫn giữa hỗ trợ và hạ thấp kỳ vọng. Nếu chỉ giao nhiệm vụ dễ hơn hoặc tách học sinh "
    "khỏi hoạt động chung, việc hỗ trợ có thể làm giảm cơ hội tham gia. Ngược lại, giữ nguyên một cách dạy cho mọi "
    "học sinh cũng không bảo đảm công bằng. Giáo viên cần giữ mục tiêu cốt lõi nhưng linh hoạt về học liệu, thời gian, "
    "phương thức tham gia và cách thể hiện kết quả."
))
add_body(doc, (
    "Công nghệ và AI làm rõ thêm mâu thuẫn này. Công cụ có thể chuyển văn bản thành giọng nói, gợi ý hình ảnh hoặc "
    "tạo học liệu nhanh hơn, nhưng đầu ra sai, ngôn ngữ định kiến và dữ liệu nhận dạng đều có thể gây hại. Vì vậy, "
    "năng lực số được đặt trong từng quyết định thiết kế, đánh giá và phối hợp. Giáo viên phải kiểm chứng nội dung, "
    "dùng dữ liệu tối thiểu cần thiết và có phương án tương đương khi học sinh không có thiết bị."
))
add_heading(doc, "3.3. Nhu cầu cần giải quyết", 2)
add_body(doc, (
    "Thực tiễn cần một công cụ đủ ngắn để giáo viên tự rà soát nhưng đủ rõ để tổ chuyên môn trao đổi bằng minh chứng. "
    "Công cụ đó phải nối được năm việc: nhận diện rào cản, thiết kế tiếp cận, tổ chức tham gia, đánh giá để điều chỉnh "
    "và phối hợp phát triển nghề nghiệp. Bên cạnh khung mô tả, giáo viên cần ví dụ về cách vận dụng trong bài học, "
    "cách chọn minh chứng và cách lập kế hoạch cải thiện mà không biến khung thành bảng chấm điểm."
))

add_heading(doc, "PHẦN 4. SẢN PHẨM ĐỀ XUẤT", 1)
add_heading(doc, "4.1. Giới thiệu sản phẩm", 2)
add_body(doc, (
    "Sản phẩm là Khung năng lực giáo dục hòa nhập của giáo viên phổ thông, được cụ thể hóa cho bối cảnh tiểu học. "
    "Khung gồm năm năng lực và chỉ số mô tả; đi kèm năm biện pháp phát triển, một tình huống lớp học, phiếu tự rà "
    "soát, kế hoạch phát triển tám tuần và hướng dẫn thu thập minh chứng."
))
add_heading(doc, "4.2. Mục tiêu và người sử dụng", 2)
add_body(doc, (
    "Khung giúp giáo viên nhận ra điểm mạnh và một ưu tiên phát triển. Tổ chuyên môn có thể xác định nội dung sinh hoạt "
    "dựa trên vấn đề thực, còn cán bộ quản lý có căn cứ chuẩn bị điều kiện hỗ trợ. Người sử dụng trực tiếp là giáo viên, tổ "
    "trưởng chuyên môn và cán bộ phụ trách giáo dục hòa nhập. Học sinh và gia đình tham gia với vai trò cung cấp tiếng "
    "nói, phản hồi và thống nhất mục tiêu hỗ trợ phù hợp."
))
add_heading(doc, "4.3. Cấu trúc và cách vận hành", 2)
add_body(doc, (
    "Năm năng lực được sắp xếp theo một chu trình nghề nghiệp nhưng không tách rời nhau. Giáo viên bắt đầu bằng việc "
    "xác định rào cản và điểm mạnh; thiết kế nhiều cách tiếp cận; tổ chức sự tham gia; thu thập minh chứng tiến bộ; "
    "sau đó phối hợp và phản tư để điều chỉnh vòng tiếp theo. Mỗi lần sử dụng chỉ nên chọn một hoặc hai chỉ số ưu "
    "tiên, tránh cộng điểm hoặc xếp hạng cá nhân."
))
add_body(doc, (
    "Trật tự này sửa một lỗi thường gặp: chọn giải pháp trước khi xác định vấn đề. Nếu chưa biết rào cản nằm ở ngôn "
    "ngữ hướng dẫn, cách thể hiện hay quan hệ trong nhóm, việc thêm thiết bị có thể không giúp ích. N1 tạo căn cứ cho "
    "N2; N2 cần N3 để trở thành sự tham gia thực tế; N4 kiểm tra tiến bộ; còn N5 giúp duy trì biện pháp vượt ra ngoài "
    "một tiết dạy. Vì vậy các bảng minh họa cùng sử dụng một hệ thống minh chứng."
))
add_heading(doc, "4.4. Khả năng áp dụng và giới hạn", 2)
add_body(doc, (
    "Khung có thể được dùng trong tự học, sinh hoạt chuyên môn, đồng thiết kế bài học và phản hồi sau dự giờ. Việc "
    "áp dụng không đòi hỏi trường phải có đầy đủ thiết bị công nghệ, vì trọng tâm là chất lượng quyết định sư phạm. "
    "Giới hạn của sản phẩm là chưa được kiểm chứng bằng dữ liệu triển khai tại một trường cụ thể; do đó các chỉ số, "
    "mốc thời gian và minh họa cần được thử nghiệm quy mô nhỏ, lấy phản hồi rồi điều chỉnh."
))

add_heading(doc, "PHẦN 5. KẾT LUẬN", 1)
add_body(doc, (
    "Giáo dục hòa nhập trong bối cảnh quốc tế hóa không thể chỉ dựa vào thiện chí. Giáo viên phải có năng lực phân "
    "tích rào cản, thiết kế linh hoạt, tổ chức sự tham gia, đánh giá vì tiến bộ và phối hợp có trách nhiệm. Khung đề "
    "xuất chuyển các yêu cầu đó thành hành vi và minh chứng gần với công việc hằng ngày ở tiểu học."
))
add_body(doc, (
    "Giá trị của sản phẩm nằm ở sự liên kết giữa khung năng lực, biện pháp phát triển và các minh họa sử dụng. Để "
    "triển khai, nhà trường cần dành thời gian sinh hoạt chuyên môn, bảo đảm nguyên tắc bảo mật, cho phép thử nghiệm "
    "trong phạm vi phù hợp và xem phản hồi của học sinh, gia đình là một nguồn minh chứng. Khung chỉ có ý nghĩa khi "
    "được dùng để hỗ trợ cải tiến, không dùng để gắn nhãn hoặc tạo thêm áp lực hành chính."
))

# PRACTICAL PRODUCT
add_heading(doc, "SẢN PHẨM THỰC HÀNH", 1, page_break=True)
title = add_body(doc, "KHUNG NĂNG LỰC GIÁO DỤC HÒA NHẬP CỦA GIÁO VIÊN PHỔ THÔNG",
                 bold=True, first_line=False, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
title.paragraph_format.keep_with_next = True

add_heading(doc, "1. Cơ sở đề xuất khung năng lực", 2)
add_heading(doc, "1.1. Xu hướng giáo dục hòa nhập và tác động của thực tiễn đối với giáo viên phổ thông", 3)
add_body(doc, (
    "Giáo dục hòa nhập chuyển từ bố trí chỗ học sang bảo đảm mọi người học được hiện diện, tham gia, tiến bộ và được "
    "tôn trọng. Cách tiếp cận dựa trên quyền, thiết kế học tập linh hoạt, đánh giá hình thành, công nghệ hỗ trợ và phối "
    "hợp đa bên làm thay đổi vai trò giáo viên: chủ động phát hiện rào cản, tạo nhiều đường tham gia và điều chỉnh "
    "từ minh chứng."
))
add_bullets(doc, [
    "Tiếp cận dựa trên quyền: tôn trọng phẩm giá, tiếng nói, quyền riêng tư và kỳ vọng học tập phù hợp.",
    "Thiết kế linh hoạt: giữ mục tiêu cốt lõi nhưng đa dạng hóa học liệu, hỗ trợ và cách thể hiện kết quả.",
    "Đánh giá vì tiến bộ: dùng thông tin trong quá trình để quyết định hỗ trợ, không gắn nhãn người học.",
    "Công nghệ tiếp cận: chọn công cụ theo mục tiêu, kiểm chứng đầu ra và có phương án không thiết bị.",
    "Hợp tác đa bên: phối hợp với học sinh, gia đình, đồng nghiệp và lực lượng hỗ trợ theo đúng vai trò.",
])
add_heading(doc, "1.2. Thực trạng năng lực giáo dục hòa nhập của giáo viên phổ thông", 3)
add_body(doc, (
    "Chưa có khảo sát tại đơn vị nên sản phẩm không đưa ra số liệu thực trạng giả định. Trước khi sử dụng khung, tổ "
    "chuyên môn cần thu thập minh chứng hiện có về nhận diện rào cản, điều chỉnh nhiệm vụ, theo dõi tiến bộ, phối hợp "
    "với gia đình và bảo vệ dữ liệu; đồng thời xác định học liệu, thiết bị, thời gian và đầu mối hỗ trợ sẵn có."
))

add_heading(doc, "2. Đề xuất khung năng lực giáo dục hòa nhập của giáo viên phổ thông", 2)
add_seq_caption(doc, "Bảng", "Khung năng lực giáo dục hòa nhập của giáo viên phổ thông")
add_table(doc, ["STT", "Năng lực", "Mô tả chỉ số năng lực"], [
    ("1", "N1. Nhận diện đa dạng và rào cản", "Thu thập thông tin đúng mục đích, không gắn nhãn; phân tích rào cản trong học liệu, nhiệm vụ, giao tiếp, môi trường và công nghệ; xác định điểm mạnh, nhu cầu và tiếng nói của học sinh trước khi chọn hỗ trợ."),
    ("2", "N2. Thiết kế dạy học tiếp cận và linh hoạt", "Xác định mục tiêu cốt lõi; tạo nhiều cách tiếp cận, tham gia và thể hiện kết quả; lựa chọn học liệu, công cụ hỗ trợ hoặc điều chỉnh hợp lý mà không hạ thấp kỳ vọng học tập."),
    ("3", "N3. Tổ chức tham gia và hỗ trợ trong lớp học", "Thiết lập an toàn tâm lý và quy tắc tôn trọng; phân nhóm, hỗ trợ bạn học và phân bổ thời gian hợp lý; xử lý hành vi, bất đồng hoặc loại trừ bằng biện pháp tích cực và bảo vệ phẩm giá."),
    ("4", "N4. Đánh giá và điều chỉnh vì sự tiến bộ", "Sử dụng quan sát, câu hỏi, sản phẩm và tự đánh giá; phản hồi cụ thể về bước tiếp theo; theo dõi mục tiêu cá nhân khi cần và điều chỉnh dạy học từ minh chứng thay vì cảm tính."),
    ("5", "N5. Phối hợp và phát triển nghề nghiệp về hòa nhập", "Phối hợp với học sinh, gia đình, đồng nghiệp và cán bộ hỗ trợ; chia sẻ tối thiểu thông tin cần thiết; tự rà soát thực hành, tham gia học tập chuyên môn và đề xuất cải thiện điều kiện hòa nhập của nhà trường."),
], [1.7, 4.5, 10.3], center_cols=(0,))
add_body(doc, "Nguồn: Tác giả đề xuất trên cơ sở UNESCO (2017), Kaplan và Bista (2022) và các quy định Việt Nam.",
         italic=True, first_line=False, before=2, after=4)

add_heading(doc, "3. Biện pháp phát triển năng lực giáo dục hòa nhập của giáo viên phổ thông", 2)
add_seq_caption(doc, "Bảng", "Biện pháp phát triển năng lực giáo dục hòa nhập của giáo viên phổ thông")
add_table(doc, ["Mục tiêu", "Cách thức thực hiện", "Nguồn lực/ phối hợp/ điều kiện"], [
    ("Xác định đúng điểm xuất phát và một ưu tiên phát triển", "Dùng khung để tự rà soát bằng minh chứng; trao đổi với một đồng nghiệp; chọn tối đa hai chỉ số cần cải thiện trong một học kỳ.", "Thời gian sinh hoạt chuyên môn; mẫu tự rà soát; nguyên tắc bảo mật và không dùng kết quả để xếp hạng."),
    ("Củng cố kiến thức về quyền, rào cản và điều chỉnh hợp lý", "Học mô-đun ngắn; phân tích tình huống; đối chiếu quy định; thực hành chuyển mô tả thiếu hụt thành phân tích rào cản và hỗ trợ.", "Tài liệu chính thống; giáo viên cốt cán; tình huống đã ẩn danh; thời gian bồi dưỡng."),
    ("Nâng năng lực thiết kế và đánh giá hòa nhập", "Đồng thiết kế một bài/chủ đề; tạo ít nhất hai cách tham gia hoặc thể hiện; dự giờ tập trung vào việc học; sửa kế hoạch từ minh chứng.", "Tổ/khối chuyên môn; chương trình, sách giáo khoa; học liệu tiếp cận được; công cụ hỗ trợ."),
    ("Tăng chất lượng phối hợp hỗ trợ học sinh", "Thống nhất mục tiêu, vai trò và kênh trao đổi với gia đình, đồng nghiệp, cán bộ hỗ trợ; chỉ chia sẻ thông tin tối thiểu cần thiết; rà tiến độ định kỳ.", "Sự đồng thuận phù hợp; đầu mối hỗ trợ; lịch phối hợp; quy trình lưu trữ và bảo vệ dữ liệu."),
    ("Duy trì cải tiến và lan tỏa thực hành có căn cứ", "Mỗi học kỳ chọn một minh chứng trước-sau; phân tích thay đổi, giới hạn và tải thực hiện; chia sẻ sản phẩm đã ẩn danh; quyết định giữ, sửa hoặc dừng biện pháp.", "Lãnh đạo chuyên môn; đồng nghiệp phản hồi; kho học liệu an toàn; tiêu chí khả thi, công bằng và tác động."),
], [4.0, 6.5, 6.0], font_size=11)
add_body(doc, "Nguồn: Tác giả đề xuất.", italic=True, first_line=False, before=2, after=4)

add_heading(doc, "4. Minh họa vận dụng khung năng lực tại tiểu học", 2)
add_heading(doc, "4.1. Tình huống minh họa", 3)
add_body(doc, (
    "Tình huống giả định diễn ra trong tiết đọc hiểu lớp 3 với mục tiêu chung: xác định thông tin chính và trình bày "
    "mối liên hệ giữa hai sự việc. Trong lớp có học sinh cần thêm gợi ý bằng hình ảnh, học sinh cần lối di chuyển rộng "
    "và một số em còn ngại phát biểu trước nhóm lớn. Giáo viên không tách các em khỏi nhiệm vụ chung mà chuẩn bị văn "
    "bản chữ rõ, thẻ trình tự, lựa chọn trả lời bằng nói, viết hoặc sắp xếp hình; đồng thời phân vai nhóm để mọi thành "
    "viên đều có đóng góp."
))
p_img = doc.add_paragraph()
p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_img.paragraph_format.keep_with_next = True
p_img.add_run().add_picture(str(ILLUSTRATION), width=Cm(14.7))
add_seq_caption(doc, "Hình", "Minh họa lớp học hòa nhập với nhiều cách tham gia", above=False)
add_body(doc, "Nguồn: Hình minh họa do AI tạo; nhân vật và tình huống đều là giả định.", italic=True,
         first_line=False, align=WD_ALIGN_PARAGRAPH.CENTER, after=5)

add_heading(doc, "4.2. Minh họa thiết kế một hoạt động học hòa nhập", 3, page_break=True)
add_seq_caption(doc, "Bảng", "Thiết kế hoạt động đọc hiểu theo hướng hòa nhập")
add_table(doc, ["Giai đoạn", "Tổ chức chung", "Điều chỉnh/hỗ trợ", "Minh chứng cần quan sát"], [
    ("Khởi động", "Nêu mục tiêu bằng lời và ghi ngắn gọn; cho cả lớp xem ba tranh gợi ý.", "Cho phép chỉ tranh, nói từ khóa hoặc trao đổi cặp đôi trước khi chia sẻ.", "Số học sinh hiểu nhiệm vụ; câu hỏi cần giải thích lại."),
    ("Khám phá văn bản", "Nhóm đọc, xác định hai sự việc và sắp xếp theo trình tự.", "Văn bản chữ rõ; thẻ sự việc; vai trò đọc, chọn thẻ, giải thích, kiểm tra.", "Mỗi thành viên có đóng góp; cách nhóm sử dụng hỗ trợ."),
    ("Thể hiện kết quả", "Nhóm trình bày mối liên hệ giữa hai sự việc.", "Chọn nói, viết đoạn ngắn, sơ đồ hoặc sắp xếp thẻ kèm giải thích.", "Mức đạt mục tiêu cốt lõi, không so sánh hình thức thể hiện."),
    ("Phản hồi", "Giáo viên nêu một điểm đã làm được và một bước tiếp theo.", "Câu phản hồi cụ thể; cho thời gian chỉnh sửa sản phẩm.", "Sản phẩm trước-sau; phản hồi của học sinh về hỗ trợ hữu ích."),
], [2.5, 4.7, 5.0, 4.3], font_size=11)
add_body(doc, (
    "Đối chiếu khung: giáo viên sử dụng N1 để xác định rào cản, N2 để tạo nhiều cách tham gia và thể hiện, N3 để "
    "phân vai, N4 để đọc sản phẩm trước-sau, và N5 khi trao đổi với đồng nghiệp hoặc gia đình về hỗ trợ cần duy trì."
), first_line=False, before=3)

add_heading(doc, "4.3. Phiếu tự rà soát minh họa", 3, page_break=True)
add_body(doc, (
    "Phiếu dưới đây minh họa cách giáo viên ghi nhận minh chứng sau một bài học. Cột tự nhận định dùng mô tả ngắn, "
    "không dùng điểm số."
))
add_seq_caption(doc, "Bảng", "Phiếu tự rà soát năng lực sau bài học minh họa")
add_table(doc, ["Năng lực", "Minh chứng hiện có", "Tự nhận định", "Ưu tiên tiếp theo"], [
    ("N1", "Đã xác định rào cản đọc chữ nhỏ và ngại nói trước lớp.", "Có căn cứ từ quan sát và phản hồi của học sinh.", "Hỏi thêm học sinh về loại gợi ý hình ảnh hữu ích."),
    ("N2", "Có ba cách thể hiện kết quả nhưng mục tiêu cốt lõi thống nhất.", "Phù hợp; cần chuẩn hóa học liệu chữ rõ.", "Chuẩn bị mẫu học liệu dùng lại cho tổ khối."),
    ("N3", "Phân vai giúp mọi thành viên tham gia; một nhóm còn để bạn làm thay.", "Chưa ổn định giữa các nhóm.", "Dạy rõ cách hỗ trợ bạn mà không làm thay."),
    ("N4", "Có sản phẩm trước-sau và phản hồi một bước tiếp theo.", "Minh chứng đủ để điều chỉnh bài kế tiếp.", "Thêm tự đánh giá ngắn của học sinh."),
    ("N5", "Đã trao đổi với đồng nghiệp về thẻ trình tự.", "Mới dừng ở chia sẻ học liệu.", "Mời đồng nghiệp dự giờ tập trung vào sự tham gia."),
], [2.0, 5.2, 4.6, 4.7], font_size=11)

add_heading(doc, "4.4. Kế hoạch phát triển năng lực trong tám tuần", 3)
add_seq_caption(doc, "Bảng", "Kế hoạch phát triển năng lực N3 và N4 trong tám tuần")
add_table(doc, ["Thời gian", "Hoạt động trọng tâm", "Người phối hợp", "Kết quả/minh chứng"], [
    ("Tuần 1", "Chọn hai chỉ số N3, N4; rà một kế hoạch bài dạy và một sản phẩm học sinh.", "Đồng nghiệp cùng khối", "Phiếu điểm xuất phát và vấn đề ưu tiên."),
    ("Tuần 2-3", "Đồng thiết kế một hoạt động có phân vai và nhiều cách thể hiện kết quả.", "Tổ chuyên môn", "Kế hoạch bài dạy, học liệu tiếp cận được."),
    ("Tuần 4", "Dạy thử; quan sát sự tham gia; thu sản phẩm trước-sau và phản hồi học sinh.", "Giáo viên dự giờ", "Phiếu quan sát, sản phẩm đã ẩn danh."),
    ("Tuần 5-6", "Điều chỉnh quy tắc hỗ trợ bạn và cách phản hồi; thực hiện lần hai.", "Học sinh, đồng nghiệp", "So sánh hai lần thực hiện."),
    ("Tuần 7-8", "Phân tích thay đổi và tải thực hiện; chia sẻ kết quả, quyết định giữ/sửa/dừng.", "Tổ trưởng chuyên môn", "Bản phản tư một trang và học liệu hoàn thiện."),
], [2.4, 7.0, 3.1, 4.0], font_size=11)

add_heading(doc, "4.5. Bộ minh chứng và cách đọc kết quả", 3, page_break=True)
add_seq_caption(doc, "Bảng", "Bộ minh chứng tối thiểu khi sử dụng khung")
add_table(doc, ["Nguồn minh chứng", "Dấu hiệu cải thiện", "Cảnh báo khi diễn giải"], [
    ("Kế hoạch và học liệu", "Mục tiêu cốt lõi rõ; có nhiều cách tiếp cận/thể hiện; hỗ trợ gắn với rào cản.", "Không kết luận chỉ từ hình thức đẹp hoặc số lượng học liệu."),
    ("Quan sát sự tham gia", "Nhiều học sinh bắt đầu nhiệm vụ, có vai trò và được trợ giúp đúng lúc.", "Không đồng nhất im lặng với không học; cần hỏi tiếng nói học sinh."),
    ("Sản phẩm trước-sau", "Sản phẩm thể hiện tiến bộ so với mục tiêu và phản hồi đã nhận.", "Không so sánh công khai giữa học sinh hoặc giữa các dạng thể hiện."),
    ("Phản hồi học sinh/gia đình", "Hỗ trợ được xem là hữu ích, tôn trọng và có thể duy trì.", "Chỉ thu thập thông tin cần thiết; bảo mật dữ liệu cá nhân."),
    ("Phản tư đồng nghiệp", "Quyết định giữ, sửa hoặc dừng biện pháp có lý do và minh chứng.", "Không biến dự giờ thành xếp hạng cá nhân."),
], [4.0, 6.5, 6.0], font_size=11)

add_heading(doc, "5. Hướng dẫn sử dụng và điều kiện triển khai", 2)
add_bullets(doc, [
    "Bước 1 - Chọn trọng tâm: giáo viên chọn một tình huống thật và tối đa hai chỉ số cần cải thiện.",
    "Bước 2 - Thu thập điểm xuất phát: dùng kế hoạch, học liệu, quan sát và tiếng nói học sinh; không tạo thêm hồ sơ nếu minh chứng hiện có đã đủ.",
    "Bước 3 - Thử nghiệm nhỏ: đồng thiết kế, thực hiện trong phạm vi được duyệt và luôn có phương án không phụ thuộc thiết bị.",
    "Bước 4 - Đọc minh chứng: so sánh trước-sau theo mục tiêu, xem cả lợi ích, tải thực hiện và tác động không mong muốn.",
    "Bước 5 - Quyết định: giữ, sửa hoặc dừng biện pháp; chia sẻ sản phẩm đã ẩn danh trong tổ chuyên môn.",
])
add_body(doc, (
    "Lưu ý sử dụng: Khung chỉ phục vụ phát triển chuyên môn. Không cộng năm năng lực thành điểm tổng, không công khai "
    "hồ sơ cá nhân và không kết luận năng lực từ một tình huống đơn lẻ. Mọi hỗ trợ phải phù hợp nhu cầu người học, "
    "nguồn lực, thẩm quyền và quy định của đơn vị."
), bold=True, first_line=False, before=5)

# REFERENCES
add_heading(doc, "TÀI LIỆU THAM KHẢO", 1, page_break=True)
references = [
    "Bộ Giáo dục và Đào tạo. (2018a). Thông tư số 03/2018/TT-BGDĐT ngày 29 tháng 01 năm 2018 quy định về giáo dục hòa nhập đối với người khuyết tật. https://vanban.chinhphu.vn/?docid=193627&pageid=27160",
    "Bộ Giáo dục và Đào tạo. (2018b). Thông tư số 32/2018/TT-BGDĐT ngày 26 tháng 12 năm 2018 ban hành Chương trình giáo dục phổ thông (được sửa đổi, bổ sung).",
    "Bộ Giáo dục và Đào tạo. (2026). Thông tư số 30/2026/TT-BGDĐT ngày 14 tháng 4 năm 2026 quy định chuẩn nghề nghiệp giáo viên cơ sở giáo dục phổ thông. https://vanban.chinhphu.vn/?classid=0&docid=217839&pageid=27160",
    "Kaplan, I., & Bista, M. B. (2022). Welcoming diversity in the learning environment: Teachers' handbook for inclusive education. UNESCO. https://unesdoc.unesco.org/ark:/48223/pf0000384009",
    "Miao, F., & Cukurova, M. (2024). AI competency framework for teachers. UNESCO. https://unesdoc.unesco.org/ark:/48223/pf0000391104",
    "OECD. (2005). Formative assessment: Improving learning in secondary classrooms. OECD Publishing. https://doi.org/10.1787/9789264007413-en",
    "Quốc hội. (2010). Luật Người khuyết tật số 51/2010/QH12 ngày 17 tháng 6 năm 2010.",
    "Redecker, C. (2017). European framework for the digital competence of educators: DigCompEdu. Publications Office of the European Union. https://doi.org/10.2760/159770",
    "UNESCO. (2017). A guide for ensuring inclusion and equity in education. https://doi.org/10.54675/MHHZ2237",
]
for ref in references:
    p = add_body(doc, ref, first_line=False, align=WD_ALIGN_PARAGRAPH.LEFT, after=4)
    p.paragraph_format.left_indent = Cm(1.0)
    p.paragraph_format.first_line_indent = Cm(-1.0)

# Enforce A4 and assignment margins on both retained sections without touching
# cover content formatting.
for section in doc.sections:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.0)

# Ask Word to refresh TOC, lists and sequence fields when the file opens.
settings = doc.settings._element
update = settings.find(qn("w:updateFields"))
if update is None:
    update = OxmlElement("w:updateFields")
    settings.append(update)
update.set(qn("w:val"), "true")

doc.core_properties.title = "Khung năng lực giáo dục hòa nhập của giáo viên phổ thông"
doc.core_properties.subject = "Tiểu luận học phần Giáo dục trong bối cảnh quốc tế hóa"
doc.core_properties.keywords = "giáo dục hòa nhập; năng lực giáo viên; tiểu học; quốc tế hóa"
doc.save(OUTPUT)
print(OUTPUT)
