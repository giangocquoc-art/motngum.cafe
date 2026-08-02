from pathlib import Path

from docx import Document
from docx.shared import Pt


BASE = Path(r"C:\Users\giang\OneDrive\Personal Vault\nghiên cứu khoa học\tiểu luận Quỳnh Phương x3\quynh-phuong-chu-de-4")
SOURCE = BASE / "output" / "K51NC_QuynhPhuong_ChuDe4.docx"
OUTPUT = BASE / "output" / "K51NC_QuynhPhuong_ChuDe4_ra-soat.docx"


def delete_paragraph(p):
    parent = p._element.getparent()
    parent.remove(p._element)
    p._p = p._element = None


def replace_exact(doc, old, new):
    matches = [p for p in doc.paragraphs if p.text.strip() == old]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one exact paragraph for {old!r}, found {len(matches)}")
    matches[0].text = new
    return matches[0]


def replace_text_everywhere(doc, old, new):
    for p in doc.paragraphs:
        if old in p.text:
            p.text = p.text.replace(old, new)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if old in p.text:
                        p.text = p.text.replace(old, new)


doc = Document(SOURCE)

# The cover already identifies the work as an end-of-course essay; the repeated
# parenthetical subtitle is redundant.
delete_paragraph(next(p for p in doc.paragraphs if p.text.strip() == "(Bài tiểu luận kết thúc học phần)"))

# Keep the learner's name only where identification/signature is expected.
full_replacements = {
    "3.1. Bối cảnh thực tiễn và điểm xuất phát của Quỳnh Phương":
        "3.1. Bối cảnh thực tiễn và phạm vi phân tích",
    "Bài viết lựa chọn xu hướng phát triển năng lực nghề nghiệp giáo viên phục vụ quốc tế hóa tại chỗ và đề xuất “La bàn năng lực giáo viên tiểu học hội nhập” gồm sáu miền, ba mức, chu trình cải tiến dựa trên minh chứng. La bàn không xếp hạng hay thay chuẩn nghề nghiệp; nó giúp xác định điểm xuất phát, chọn ưu tiên và kiểm tra tiến bộ. Sản phẩm được cá nhân hóa từ thiết kế đánh giá thực của Quỳnh Phương cho chủ điểm “Thiên nhiên kì thú” ở Tiếng Việt lớp 5. Đây là điểm xuất phát thiết kế, không phải kết quả đánh giá tại một trường.":
        "Bài viết lựa chọn xu hướng phát triển năng lực nghề nghiệp giáo viên phục vụ quốc tế hóa tại chỗ và đề xuất “La bàn năng lực giáo viên tiểu học hội nhập” gồm sáu miền, ba mức và chu trình cải tiến dựa trên minh chứng. La bàn không xếp hạng hoặc thay thế chuẩn nghề nghiệp; công cụ giúp giáo viên xác định điểm xuất phát, chọn ưu tiên và kiểm tra tiến bộ. Sản phẩm được phát triển từ một thiết kế đánh giá thực cho chủ điểm “Thiên nhiên kì thú” ở Tiếng Việt lớp 5. Đây chỉ là điểm xuất phát thiết kế, không phải kết quả đánh giá năng lực tại một trường.",
    "Hồ sơ hiện có không nêu tên đơn vị công tác hay cung cấp dữ liệu khảo sát cấp trường. Vì vậy, bài viết không gán cho một trường học ở Đà Nẵng những thực trạng chưa được kiểm chứng. Phạm vi thực tiễn được xác định ở cấp cá nhân: sản phẩm Quỳnh Phương đã xây dựng trong quá trình đào tạo và tình huống áp dụng đề xuất ở lớp 5; những nhận định về triển khai rộng hơn chỉ được trình bày dưới dạng nhu cầu, điều kiện và kế hoạch thử nghiệm.":
        "Do chưa có dữ liệu khảo sát tại một đơn vị cụ thể, bài viết không suy rộng thành nhận định hoặc số liệu cấp trường. Phần thực tiễn tập trung vào một tình huống áp dụng dự kiến ở lớp 5 và những điều kiện tối thiểu cần kiểm chứng trước khi triển khai. Các nhận định về cơ hội, thách thức được dùng để thiết kế tiêu chí thử nghiệm, không phải kết luận về thực trạng của một cơ sở giáo dục.",
    "Hồ sơ dự án cho thấy Quỳnh Phương từng xây dựng phương án đánh giá thực cho chủ điểm “Thiên nhiên kì thú”, bài 9–16 của Tiếng Việt 5, tập một, bộ Kết nối tri thức với cuộc sống. Thiết kế có nhiệm vụ “Hướng dẫn viên trẻ vì thiên nhiên”, rubric, minh chứng cá nhân–nhóm, phản hồi và phương án không thiết bị. Điểm xuất phát phù hợp vì thế là thiết kế học tập và đánh giá dựa trên minh chứng, không phải thao tác công nghệ đơn lẻ.":
        "Điểm xuất phát thiết kế là một phương án đánh giá thực cho chủ điểm “Thiên nhiên kì thú”, bài 9–16 của Tiếng Việt 5, tập một, bộ Kết nối tri thức với cuộc sống. Phương án có nhiệm vụ “Hướng dẫn viên trẻ vì thiên nhiên”, bảng tiêu chí, minh chứng cá nhân–nhóm, phản hồi và lựa chọn không dùng thiết bị. Kinh nghiệm phù hợp để phát triển tiếp là thiết kế học tập và đánh giá dựa trên minh chứng, không phải thao tác công nghệ đơn lẻ.",
    "Tiếng Việt lớp 5 cho phép tích hợp quốc tế hóa mà không tạo môn riêng. Học sinh có thể đọc văn bản từ nhiều bối cảnh, so sánh cách diễn đạt quan hệ với thiên nhiên, kiểm tra nguồn ảnh, giới thiệu cảnh quan địa phương và phản tư về trách nhiệm. Kinh nghiệm thiết kế nhiệm vụ–rubric giúp Quỳnh Phương chuyển các ý tưởng này thành minh chứng.":
        "Tiếng Việt lớp 5 cho phép tích hợp quốc tế hóa mà không tạo môn học riêng. Học sinh có thể đọc văn bản từ nhiều bối cảnh, so sánh cách diễn đạt mối quan hệ với thiên nhiên, kiểm tra nguồn ảnh, giới thiệu cảnh quan địa phương và phản tư về trách nhiệm. Kinh nghiệm thiết kế nhiệm vụ và bảng tiêu chí là điểm tựa để chuyển các ý tưởng này thành minh chứng.",
    "Với Quỳnh Phương, M2 và M4 là nền từ kinh nghiệm đánh giá thực ở Tiếng Việt lớp 5; M1 và M3 là hướng mở bằng rà nhiều góc nhìn, bản sắc và khả năng tiếp cận. M5 phát triển qua kiểm chứng nguồn/AI và bảo vệ dữ liệu; M6 chỉ sang Lan tỏa sau ít nhất hai chu kỳ có phản hồi đồng nghiệp.":
        "Từ minh chứng nghề nghiệp hiện có, M2 và M4 có thể được xem là nền ban đầu; M1 và M3 là hướng ưu tiên để mở rộng nhiều góc nhìn, bản sắc và khả năng tiếp cận. M5 cần phát triển qua kiểm chứng nguồn, đầu ra AI và bảo vệ dữ liệu; M6 chỉ chuyển sang mức Lan tỏa sau ít nhất hai chu kỳ có phản hồi đồng nghiệp.",
    "Đối với Quỳnh Phương, sản phẩm không phủ nhận kinh nghiệm đã có mà dùng chính nền tảng thiết kế đánh giá thực ở Tiếng Việt lớp 5 làm điểm khởi hành. Giá trị quan trọng nhất là tạo một con đường phát triển có lựa chọn: không chạy theo mọi xu hướng, không dùng công nghệ để thay mục tiêu và không thu gom hồ sơ tách khỏi việc học của trẻ.":
        "Sản phẩm sử dụng nền tảng thiết kế đánh giá thực ở Tiếng Việt lớp 5 làm điểm khởi hành, đồng thời chỉ rõ những miền còn cần kiểm chứng. Giá trị quan trọng nhất là tạo một con đường phát triển có lựa chọn: không chạy theo mọi xu hướng, không dùng công nghệ để thay mục tiêu và không thu gom hồ sơ tách khỏi việc học của học sinh.",
    "Điểm xuất phát cá nhân hóa: Hồ sơ chuyên môn trong dự án cho thấy người thiết kế đã phát triển phương án đánh giá thực cho môn Tiếng Việt lớp 5, có kinh nghiệm căn chỉnh yêu cầu cần đạt–nhiệm vụ–minh chứng–rubric. Thông tin này dùng để chọn hướng phát triển, không thay cho đánh giá năng lực tại đơn vị công tác.":
        "Điểm xuất phát minh họa: Người viết đã có một phương án đánh giá thực cho Tiếng Việt lớp 5, trong đó căn chỉnh yêu cầu cần đạt–nhiệm vụ–minh chứng–bảng tiêu chí. Kinh nghiệm này chỉ được xem là minh chứng ứng viên để chọn hướng phát triển, không được dùng như kết luận năng lực tại đơn vị công tác.",
    "2.2.2. Định vị ban đầu gợi ý cho Quỳnh Phương":
        "2.2.2. Ví dụ định vị ban đầu từ hồ sơ chuyên môn",
    "2.1. Rubric ba mức phát triển":
        "2.1. Bảng mô tả ba mức phát triển",
    "Bảng 4. Rubric ba mức phát triển theo sáu miền":
        "Bảng 4. Ba mức phát triển theo sáu miền",
    "Bảng 6. Kế hoạch 90 ngày gợi ý cho Quỳnh Phương":
        "Bảng 6. Kế hoạch 90 ngày minh họa",
    "Bảng 8. Kế hoạch phát triển năng lực nghề nghiệp cá nhân của Quỳnh Phương giai đoạn 2026–2031":
        "Bảng 8. Kế hoạch phát triển năng lực nghề nghiệp cá nhân giai đoạn 2026–2031",
}

for old, new in full_replacements.items():
    # Some strings also occur in the static TOC/list. Replace every exact
    # paragraph match while retaining each paragraph's existing style.
    matches = [p for p in doc.paragraphs if p.text.strip() == old]
    if not matches:
        raise RuntimeError(f"Missing expected paragraph: {old!r}")
    for p in matches:
        p.text = new

# Static TOC/list entries include page numbers, so handle their full strings.
toc_and_list = {
    "3.1. Bối cảnh thực tiễn và điểm xuất phát của Quỳnh Phương\t4":
        "3.1. Bối cảnh thực tiễn và phạm vi phân tích\t4",
    "Bảng 6. Kế hoạch 90 ngày gợi ý cho Quỳnh Phương\t15":
        "Bảng 6. Kế hoạch 90 ngày minh họa\t15",
    "Bảng 8. Kế hoạch phát triển năng lực nghề nghiệp cá nhân của Quỳnh Phương giai đoạn 2026–2031\t17":
        "Bảng 8. Kế hoạch phát triển năng lực nghề nghiệp cá nhân giai đoạn 2026–2031\t17",
    "Bảng 4. Rubric ba mức phát triển theo sáu miền\t12":
        "Bảng 4. Ba mức phát triển theo sáu miền\t13",
}
for old, new in toc_and_list.items():
    replace_exact(doc, old, new)

# The author is already identified on the cover and signs the acknowledgments.
# A repeated designer credit inside the product is unnecessary.
delete_paragraph(next(p for p in doc.paragraphs if p.text.strip() == "Người thiết kế: Quỳnh Phương."))

# Remove avoidable English jargon and improve a few mechanically awkward terms.
for old, new in [
    ("agency của người học", "quyền chủ động của người học"),
    ("teacher agency", "quyền chủ động nghề nghiệp"),
    ("rubric ba mức", "bảng mô tả ba mức"),
    ("rubric", "bảng tiêu chí"),
    ("protocol kết nối", "quy trình kết nối"),
    ("một bảng khung năng lực", "một bản khung năng lực"),
    ("đối với phát triển nghề nghiệp", "đối với việc phát triển nghề nghiệp"),
    ("dựa trên độ sử dụng", "dựa trên dữ liệu sử dụng"),
]:
    replace_text_everywhere(doc, old, new)

# Separate the academic references from the product's final evaluation table.
refs = next(p for p in doc.paragraphs if p.text.strip() == "TÀI LIỆU THAM KHẢO")
refs.paragraph_format.page_break_before = True

# Match the teacher's stated typography. Headings retain their hierarchy; body,
# captions, sources, TOC and all table text inherit 12 pt Times New Roman.
for style_name in ["Normal", "Figure Caption", "Table Caption", "Table Source", "toc 1", "toc 2"]:
    style = doc.styles[style_name]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

doc.core_properties.last_modified_by = "Quỳnh Phương"
doc.save(OUTPUT)
print(OUTPUT)
