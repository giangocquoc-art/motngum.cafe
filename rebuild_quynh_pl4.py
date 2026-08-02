from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


BASE = Path(r"C:\Users\giang\OneDrive\Personal Vault\nghiên cứu khoa học\tiểu luận Quỳnh Phương x3\quynh-phuong-chu-de-4")
SOURCE = BASE / "output" / "K51NC_QuynhPhuong_ChuDe4_ra-soat.docx"
OUTPUT = BASE / "output" / "K51NC_QuynhPhuong_ChuDe4_PL4.docx"


def remove_between(body, start_el, end_el, include_start=False):
    children = list(body)
    a = children.index(start_el) + (0 if include_start else 1)
    b = children.index(end_el)
    for el in children[a:b]:
        body.remove(el)


def add_tabbed_before(anchor, text, style="toc 1", indent_cm=0):
    p = anchor.insert_paragraph_before(style=style)
    p.paragraph_format.left_indent = Cm(indent_cm)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(16.2), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    p.add_run(text)
    return p


def add_para(doc, text="", style="Normal", bold=False, italic=False, center=False, before=0, after=4):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    return p


def add_bullets(doc, items):
    for item in items:
        p = add_para(doc, item, "List Bullet", after=2)
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.3)


def shade(cell, fill="E7E6E6"):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def repeat_header(row):
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(marker)


def set_table_widths(table, widths_cm):
    widths = [round(x * 567) for x in widths_cm]
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tcw = cell._tc.get_or_add_tcPr().tcW
            tcw.set(qn("w:w"), str(widths[i]))
            tcw.set(qn("w:type"), "dxa")


def fill_cell(cell, text, bold=False, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.1
    p.paragraph_format.left_indent = Cm(0)
    p.paragraph_format.right_indent = Cm(0)
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_pl4_table(doc, headers, rows, widths_cm, center_cols=()):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, header in enumerate(headers):
        shade(table.rows[0].cells[i])
        fill_cell(table.rows[0].cells[i], header, bold=True, center=True)
        if header == "STT":
            table.rows[0].cells[i]._tc.get_or_add_tcPr().append(OxmlElement("w:noWrap"))
    repeat_header(table.rows[0])
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
    for values in rows:
        row = table.add_row()
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
        cells = row.cells
        for i, value in enumerate(values):
            fill_cell(cells[i], value, center=i in center_cols)
    set_table_widths(table, widths_cm)
    return table


doc = Document(SOURCE)
body = doc._element.body

# Cover and acknowledgments: align the identity of the report with PL4.
for p in doc.paragraphs:
    text = p.text.strip()
    if text.startswith("ĐỀ TÀI:"):
        p.text = (
            "ĐỀ TÀI: KHUNG NĂNG LỰC GIÁO DỤC HÒA NHẬP CỦA GIÁO VIÊN PHỔ THÔNG – "
            "VẬN DỤNG Ở TIỂU HỌC TRONG BỐI CẢNH QUỐC TẾ HÓA GIÁO DỤC"
        )
    elif text.startswith("Tôi trân trọng cảm ơn những đồng nghiệp"):
        p.text = (
            "Tôi trân trọng cảm ơn các đồng nghiệp đã chia sẻ những tình huống thực tế về sự khác biệt "
            "giữa học sinh, rào cản tham gia và nhu cầu phối hợp với gia đình. Những trao đổi ấy giúp tôi "
            "chuyển yêu cầu của mẫu PL4 thành một khung năng lực có chỉ số quan sát được và các biện pháp "
            "phát triển gắn với công việc hằng ngày của giáo viên tiểu học."
        )
    elif text.startswith("Bài tiểu luận được hoàn thành trên cơ sở"):
        p.text = (
            "Bài tiểu luận được hoàn thành trên cơ sở nghiên cứu tài liệu và thiết kế sư phạm, chưa thay thế "
            "khảo sát năng lực tại một đơn vị cụ thể và chưa chứng minh hiệu quả bằng dữ liệu triển khai. "
            "Tôi mong tiếp tục nhận được góp ý về độ rõ của chỉ số, tính khả thi của biện pháp và cách bảo vệ "
            "phẩm giá, dữ liệu, quyền tham gia của học sinh khi vận dụng khung năng lực."
        )

# Rebuild front matter to remove all PL3-only entries and the obsolete figure list.
toc_title = next(p for p in doc.paragraphs if p.text.strip() == "MỤC LỤC")
table_list_title = next(p for p in doc.paragraphs if p.text.strip() == "DANH MỤC BẢNG")
section_break_p = next(
    p for p in doc.paragraphs
    if p._p.pPr is not None and p._p.pPr.sectPr is not None
)
remove_between(body, toc_title._p, table_list_title._p)
remove_between(body, table_list_title._p, section_break_p._p)

toc_entries = [
    ("PHẦN 1. GIỚI THIỆU VẤN ĐỀ\t1", "toc 1", 0),
    ("PHẦN 2. CƠ SỞ LÝ LUẬN\t1", "toc 1", 0),
    ("2.1. Các khái niệm liên quan\t1", "toc 2", 0.6),
    ("2.2. Giáo dục hòa nhập trong bối cảnh quốc tế hóa\t2", "toc 2", 0.6),
    ("2.3. Nguyên tắc thiết kế khung năng lực\t2", "toc 2", 0.6),
    ("PHẦN 3. THỰC TIỄN VÀ NHU CẦU\t2", "toc 1", 0),
    ("3.1. Phạm vi thực tiễn\t2", "toc 2", 0.6),
    ("3.2. Cơ hội\t2", "toc 2", 0.6),
    ("3.3. Thách thức và nhu cầu cần giải quyết\t3", "toc 2", 0.6),
    ("PHẦN 4. SẢN PHẨM ĐỀ XUẤT\t3", "toc 1", 0),
    ("PHẦN 5. KẾT LUẬN\t4", "toc 1", 0),
    ("SẢN PHẨM THỰC HÀNH\t5", "toc 1", 0),
    ("1. Cơ sở đề xuất khung năng lực\t5", "toc 2", 0.6),
    ("2. Đề xuất khung năng lực giáo dục hòa nhập của giáo viên phổ thông\t5", "toc 2", 0.6),
    ("3. Biện pháp phát triển năng lực giáo dục hòa nhập của giáo viên phổ thông\t6", "toc 2", 0.6),
    ("TÀI LIỆU THAM KHẢO\t8", "toc 1", 0),
]
for text, style, indent in toc_entries:
    add_tabbed_before(table_list_title, text, style, indent)
pb = table_list_title.insert_paragraph_before()
pb.add_run().add_break(WD_BREAK.PAGE)

add_tabbed_before(section_break_p, "Bảng 1. Khung năng lực giáo dục hòa nhập của giáo viên phổ thông\t5", "Normal", 0)
add_tabbed_before(section_break_p, "Bảng 2. Biện pháp phát triển năng lực giáo dục hòa nhập của giáo viên phổ thông\t6", "Normal", 0)

# Replace the old report and practical product as one coherent PL4 submission.
old_report_start = next(p for p in doc.paragraphs if p.text.strip() == "PHẦN 1. GIỚI THIỆU VẤN ĐỀ")
refs = next(p for p in doc.paragraphs if p.text.strip() == "TÀI LIỆU THAM KHẢO")
remove_between(body, old_report_start._p, refs._p, include_start=True)
start = len(list(body)) - 1

add_para(doc, "PHẦN 1. GIỚI THIỆU VẤN ĐỀ", "Heading 1")
add_para(doc, (
    "Giáo dục hòa nhập không chỉ là sắp xếp cho học sinh khuyết tật học cùng lớp với các bạn. Bản chất của "
    "hòa nhập là nhận diện và giảm những rào cản khiến bất kỳ học sinh nào khó tiếp cận, tham gia, tiến bộ "
    "hoặc cảm thấy mình không thuộc về lớp học. Rào cản có thể xuất phát từ học liệu, cách giao nhiệm vụ, "
    "giao tiếp, môi trường vật chất, công nghệ, định kiến hoặc cách đánh giá quá cứng nhắc."
))
add_para(doc, (
    "Trong bối cảnh quốc tế hóa, lớp học ngày càng tiếp xúc với nhiều ngôn ngữ, nguồn học liệu, công nghệ "
    "và quan niệm khác nhau về sự đa dạng. Quốc tế hóa chỉ có ý nghĩa khi mở rộng cơ hội học tập thay vì tạo "
    "thêm chênh lệch. UNESCO (2017) đặt hòa nhập và công bằng ở trung tâm của chất lượng giáo dục; pháp luật "
    "Việt Nam cũng xác định giáo dục hòa nhập là phương thức chủ yếu đối với người khuyết tật (Quốc hội, 2010; "
    "Bộ Giáo dục và Đào tạo, 2018a)."
))
add_para(doc, (
    "Bài viết vì vậy lựa chọn xây dựng Khung năng lực giáo dục hòa nhập của giáo viên phổ thông, vận dụng ở "
    "tiểu học. Sản phẩm bám đúng PL4: cơ sở đề xuất; khung năng lực với chỉ số quan sát được; và biện pháp phát "
    "triển năng lực. Khung không dùng để gắn nhãn hoặc xếp hạng giáo viên, mà giúp tự rà soát, chọn ưu tiên bồi "
    "dưỡng và cải tiến thực hành dựa trên minh chứng."
))

add_para(doc, "PHẦN 2. CƠ SỞ LÝ LUẬN", "Heading 1", before=8)
add_para(doc, "2.1. Các khái niệm liên quan", "Heading 2")
add_para(doc, (
    "Giáo dục hòa nhập là quá trình tăng năng lực của hệ thống giáo dục để tiếp cận mọi người học, coi sự đa "
    "dạng là điều kiện bình thường của lớp học chứ không phải vấn đề của riêng một nhóm. Cách tiếp cận này dịch "
    "chuyển câu hỏi từ “học sinh thiếu gì” sang “rào cản nào đang cản trở học sinh và môi trường cần thay đổi "
    "ra sao”. Hòa nhập vì thế liên quan đồng thời tới sự hiện diện, tham gia, tiến bộ và cảm giác thuộc về."
))
add_para(doc, (
    "Năng lực giáo dục hòa nhập của giáo viên là khả năng huy động kiến thức, kỹ năng, thái độ và giá trị để "
    "nhận diện nhu cầu, thiết kế dạy học linh hoạt, tổ chức hỗ trợ, đánh giá vì sự tiến bộ và phối hợp với các "
    "bên liên quan. Năng lực phải được thể hiện qua hành vi và minh chứng nghề nghiệp, không chỉ qua chứng chỉ "
    "hoặc sự tự tin chủ quan."
))
add_para(doc, (
    "Điều chỉnh hợp lý là thay đổi cần thiết và phù hợp để người học có thể tham gia mà không làm mất mục tiêu "
    "cốt lõi. Phân hóa và thiết kế học tập linh hoạt tạo nhiều cách tiếp cận, tham gia và thể hiện kết quả. "
    "Đánh giá vì sự tiến bộ dùng quan sát, câu hỏi, sản phẩm và phản hồi để quyết định bước hỗ trợ tiếp theo "
    "(OECD, 2005)."
))

add_para(doc, "2.2. Giáo dục hòa nhập trong bối cảnh quốc tế hóa", "Heading 2")
add_para(doc, (
    "Xu hướng thứ nhất là chuyển từ tiếp cận theo thiếu hụt sang tiếp cận dựa trên quyền, công bằng và sự tham "
    "gia. Giáo viên không được biến chẩn đoán thành giới hạn cố định, không yêu cầu một học sinh đại diện cho "
    "cả nhóm và không hạ thấp kỳ vọng chỉ vì người học cần hỗ trợ. Mọi quyết định phải tôn trọng phẩm giá, tiếng "
    "nói và quyền riêng tư của học sinh."
))
add_para(doc, (
    "Xu hướng thứ hai là thiết kế học tập linh hoạt và tăng khả năng tiếp cận. Học liệu số, công cụ hỗ trợ và AI "
    "có thể chuyển văn bản thành giọng nói, tạo phương án biểu đạt hoặc hỗ trợ giáo viên điều chỉnh nhiệm vụ. "
    "Tuy nhiên, công nghệ chỉ có giá trị khi phù hợp mục tiêu, có phương án không thiết bị, được kiểm chứng và "
    "không làm lộ dữ liệu cá nhân (Redecker, 2017; Miao & Cukurova, 2024)."
))
add_para(doc, (
    "Xu hướng thứ ba là hợp tác đa bên và học tập nghề nghiệp tại chỗ. Giáo viên chủ nhiệm không thể một mình "
    "giải quyết mọi nhu cầu; cần phối hợp với học sinh, gia đình, đồng nghiệp, cán bộ hỗ trợ và lãnh đạo. Việc "
    "chia sẻ thông tin phải theo nguyên tắc tối thiểu cần thiết, đúng mục đích và có sự đồng thuận phù hợp. Phát "
    "triển năng lực nên diễn ra qua phân tích tình huống, đồng thiết kế, quan sát việc học và sửa kế hoạch."
))

add_para(doc, "2.3. Nguyên tắc thiết kế khung năng lực", "Heading 2")
add_para(doc, (
    "Khung được xây dựng theo năm nguyên tắc: bám nhiệm vụ thực của giáo viên phổ thông; dùng chỉ số có thể quan "
    "sát; xem rào cản trong môi trường là đối tượng cần thay đổi; giữ mục tiêu học tập nhưng linh hoạt con đường "
    "tham gia; và dùng minh chứng để cải tiến thay vì xếp hạng. Chỉ số phải đủ rõ để hai người đọc hiểu tương đối "
    "thống nhất, đồng thời đủ mở để điều chỉnh theo dạng nhu cầu và điều kiện nhà trường."
))
add_para(doc, (
    "Khung cũng phải tương thích với Chương trình giáo dục phổ thông 2018, quy định về giáo dục hòa nhập đối với "
    "người khuyết tật và chuẩn nghề nghiệp giáo viên hiện hành (Bộ Giáo dục và Đào tạo, 2018a, 2018b, 2026). "
    "Vì vậy, sản phẩm không tạo thêm một chuẩn pháp lý mà cụ thể hóa những hành vi nghề nghiệp cần được bồi dưỡng."
))

add_para(doc, "PHẦN 3. THỰC TIỄN VÀ NHU CẦU", "Heading 1", before=8)
add_para(doc, "3.1. Phạm vi thực tiễn", "Heading 2")
add_para(doc, (
    "Bài viết chưa có dữ liệu khảo sát tại một trường cụ thể nên không đưa ra tỷ lệ học sinh cần hỗ trợ, mức năng "
    "lực của giáo viên hoặc kết quả can thiệp giả định. Phân tích thực tiễn tập trung vào những quyết định thường "
    "gặp ở tiểu học: nhận biết rào cản, điều chỉnh nhiệm vụ, tạo nhiều cách tham gia, phản hồi vì tiến bộ, phối hợp "
    "với gia đình và bảo vệ thông tin nhạy cảm. Các nội dung này phải được kiểm chứng khi thử nghiệm tại đơn vị."
))
add_para(doc, "3.2. Cơ hội", "Heading 2")
add_para(doc, (
    "Chương trình định hướng phẩm chất, năng lực tạo không gian cho giáo viên sử dụng nhiệm vụ đa dạng và đánh giá "
    "thường xuyên. Sinh hoạt tổ chuyên môn có thể trở thành nơi phân tích rào cản, đồng thiết kế học liệu và xem "
    "minh chứng học tập. Tài nguyên số và học liệu mở giúp tạo bản chữ lớn, âm thanh, hình ảnh, từ khóa hoặc lựa "
    "chọn biểu đạt nếu được chọn lọc và kiểm tra khả năng tiếp cận."
))
add_para(doc, "3.3. Thách thức và nhu cầu cần giải quyết", "Heading 2")
add_para(doc, (
    "Khó khăn phổ biến là đồng nhất hòa nhập với chăm sóc riêng học sinh khuyết tật, hoặc giao toàn bộ trách nhiệm "
    "cho giáo viên chủ nhiệm. Khi thiếu chỉ số rõ, giáo viên dễ chọn giải pháp hình thức như giảm yêu cầu, cho làm "
    "bài khác hoàn toàn hoặc thu nhiều hồ sơ nhưng không thay đổi cách dạy. Việc sử dụng thông tin chẩn đoán, hình "
    "ảnh và bài làm cũng có nguy cơ xâm phạm riêng tư nếu không giới hạn mục đích và người tiếp cận."
))
add_para(doc, (
    "Nhu cầu cốt lõi là một khung ngắn gọn liên kết trực tiếp năng lực với chỉ số, sau đó nối mỗi nhóm năng lực "
    "với biện pháp phát triển, nguồn lực và điều kiện. Khung cần hỗ trợ giáo viên nhận ra điểm mạnh, chọn một ưu "
    "tiên khả thi và kiểm tra thay đổi qua kế hoạch bài dạy, học liệu đã điều chỉnh, quan sát tham gia, phản hồi "
    "của học sinh và trao đổi với gia đình hoặc đồng nghiệp."
))

add_para(doc, "PHẦN 4. SẢN PHẨM ĐỀ XUẤT", "Heading 1", before=8)
add_para(doc, "4.1. Cấu trúc và mục tiêu", "Heading 2")
add_para(doc, (
    "Sản phẩm là Khung năng lực giáo dục hòa nhập của giáo viên phổ thông, vận dụng ở tiểu học, được trình bày "
    "đúng ba phần của PL4. Phần cơ sở giải thích xu hướng và phạm vi thực tiễn; phần khung gồm năm năng lực với "
    "các chỉ số quan sát được; phần biện pháp nối mục tiêu bồi dưỡng với cách thực hiện, nguồn lực và điều kiện."
))
add_para(doc, "4.2. Nội dung khung năng lực", "Heading 2")
add_para(doc, (
    "Năm năng lực tạo thành một chuỗi quyết định: nhận diện đa dạng và rào cản; thiết kế dạy học tiếp cận; tổ chức "
    "tham gia và hỗ trợ; đánh giá, điều chỉnh theo tiến bộ cá nhân; phối hợp và phát triển nghề nghiệp. Mỗi năng lực "
    "được mô tả bằng ba nhóm hành vi để tránh các nhãn chung như “quan tâm học sinh” hoặc “có tinh thần hòa nhập”."
))
add_para(doc, "4.3. Khả năng áp dụng", "Heading 2")
add_para(doc, (
    "Khung có thể dùng trong tự rà soát, xây kế hoạch bồi dưỡng tổ chuyên môn hoặc thiết kế một chu kỳ nghiên cứu "
    "bài học. Khi áp dụng, giáo viên chỉ chọn một hoặc hai chỉ số ưu tiên, xác định minh chứng và hỗ trợ cần thiết. "
    "Nhà trường không cộng các năng lực thành điểm thi đua và không công khai hồ sơ cá nhân. Hiệu quả chỉ được "
    "kết luận sau khi có dữ liệu về mức tham gia, tiến bộ, tính khả thi và phản hồi của người học."
))

pb = doc.add_paragraph()
pb.add_run().add_break(WD_BREAK.PAGE)
add_para(doc, "PHẦN 5. KẾT LUẬN", "Heading 1")
add_para(doc, (
    "Giáo dục hòa nhập trong bối cảnh quốc tế hóa đòi hỏi giáo viên vừa bảo đảm quyền và công bằng, vừa có khả "
    "năng thiết kế, tương tác, đánh giá, phối hợp và sử dụng công nghệ có trách nhiệm. Khung đề xuất chuyển yêu "
    "cầu rộng thành năm năng lực và các chỉ số có thể quan sát, nhờ đó giáo viên có điểm tựa để chọn ưu tiên phát "
    "triển thay vì xử lý tình huống rời rạc."
))
add_para(doc, (
    "Giá trị của sản phẩm nằm ở sự bám sát PL4, tính gọn và khả năng nối trực tiếp với biện pháp bồi dưỡng. Điều "
    "kiện triển khai gồm thời gian sinh hoạt chuyên môn, quyền thử nghiệm trong phạm vi chương trình, sự phối hợp "
    "của gia đình và cán bộ hỗ trợ, học liệu tiếp cận được, quy trình bảo vệ dữ liệu và cam kết không dùng khung "
    "để gắn nhãn giáo viên hoặc học sinh."
))

pb = doc.add_paragraph()
pb.add_run().add_break(WD_BREAK.PAGE)
add_para(doc, "SẢN PHẨM THỰC HÀNH", "Heading 1", center=True)
add_para(doc, "KHUNG NĂNG LỰC GIÁO DỤC HÒA NHẬP CỦA GIÁO VIÊN PHỔ THÔNG", "Heading 2", center=True, after=8)

add_para(doc, "1. Cơ sở đề xuất khung năng lực", "Heading 2")
add_para(doc, "1.1. Xu hướng giáo dục hòa nhập và tác động của thực tiễn đối với giáo viên phổ thông", "Heading 3", italic=True)
add_para(doc, (
    "Giáo dục hòa nhập chuyển từ việc bố trí chỗ học sang bảo đảm mọi người học được hiện diện, tham gia, tiến bộ "
    "và được tôn trọng. Cách tiếp cận dựa trên quyền, thiết kế học tập linh hoạt, công nghệ hỗ trợ và phối hợp đa "
    "bên làm thay đổi vai trò giáo viên: từ truyền đạt cùng một cách sang chủ động phát hiện rào cản, thiết kế "
    "nhiều đường tham gia và điều chỉnh từ minh chứng."
))
add_bullets(doc, [
    "Tiếp cận dựa trên quyền: tôn trọng phẩm giá, tiếng nói, quyền riêng tư và kỳ vọng học tập phù hợp.",
    "Thiết kế linh hoạt: giữ mục tiêu cốt lõi nhưng đa dạng hóa học liệu, hỗ trợ và cách thể hiện kết quả.",
    "Đánh giá vì tiến bộ: dùng thông tin trong quá trình để quyết định hỗ trợ, không gắn nhãn người học.",
    "Công nghệ tiếp cận: chọn công cụ theo mục tiêu, kiểm chứng đầu ra và luôn có phương án không thiết bị.",
    "Hợp tác đa bên: phối hợp với học sinh, gia đình, đồng nghiệp và lực lượng hỗ trợ theo đúng vai trò."
])
add_para(doc, "1.2. Thực trạng năng lực giáo dục hòa nhập của giáo viên phổ thông", "Heading 3", italic=True)
add_para(doc, (
    "Chưa có khảo sát tại đơn vị nên sản phẩm không đưa ra số liệu thực trạng giả định. Từ yêu cầu nghề nghiệp và "
    "những tình huống thường gặp, có thể xác định các khoảng trống cần được khảo sát: mức hiểu về hòa nhập; khả "
    "năng nhận diện rào cản; chất lượng điều chỉnh nhiệm vụ; cách theo dõi tiến bộ; sự phối hợp với gia đình; và "
    "việc bảo vệ dữ liệu. Trước khi sử dụng khung, tổ chuyên môn cần thu thập minh chứng hiện có, nhu cầu hỗ trợ, "
    "điều kiện học liệu–thiết bị và đầu mối phối hợp."
))

add_para(doc, "2. Đề xuất khung năng lực giáo dục hòa nhập của giáo viên phổ thông", "Heading 2", before=8)
add_para(doc, "Bảng 1. Khung năng lực giáo dục hòa nhập của giáo viên phổ thông", "Table Caption", center=True)
add_pl4_table(doc,
    ["STT", "Năng lực", "Mô tả chỉ số năng lực"],
    [
        ("1", "N1. Nhận diện đa dạng và rào cản", "Thu thập thông tin đúng mục đích, không gắn nhãn; phân tích rào cản trong học liệu, nhiệm vụ, giao tiếp, môi trường và công nghệ; xác định điểm mạnh, nhu cầu, tiếng nói của học sinh trước khi chọn hỗ trợ."),
        ("2", "N2. Thiết kế dạy học tiếp cận và linh hoạt", "Xác định mục tiêu cốt lõi; tạo nhiều cách tiếp cận, tham gia và thể hiện kết quả; lựa chọn học liệu, công cụ hỗ trợ hoặc điều chỉnh hợp lý mà không hạ thấp kỳ vọng học tập."),
        ("3", "N3. Tổ chức tham gia và hỗ trợ trong lớp học", "Thiết lập an toàn tâm lý và quy tắc tôn trọng; phân nhóm, hỗ trợ bạn học và phân bổ thời gian hợp lý; xử lý hành vi, bất đồng hoặc loại trừ bằng biện pháp tích cực và bảo vệ phẩm giá."),
        ("4", "N4. Đánh giá và điều chỉnh vì sự tiến bộ", "Sử dụng quan sát, câu hỏi, sản phẩm và tự đánh giá; phản hồi cụ thể về bước tiếp theo; theo dõi mục tiêu cá nhân khi cần và điều chỉnh dạy học từ minh chứng thay vì từ cảm tính."),
        ("5", "N5. Phối hợp và phát triển nghề nghiệp về hòa nhập", "Phối hợp với học sinh, gia đình, đồng nghiệp và cán bộ hỗ trợ; chia sẻ tối thiểu thông tin cần thiết; tự rà soát thực hành, tham gia học tập chuyên môn và đề xuất cải thiện điều kiện hòa nhập của nhà trường."),
    ],
    [2.0, 4.5, 10.0],
    center_cols=(0,)
)
add_para(doc, "Nguồn: Tác giả đề xuất trên cơ sở UNESCO (2017), Kaplan và Bista (2022), quy định Việt Nam và các khung năng lực số–AI dành cho giáo viên.", "Table Source", italic=True)

measures_heading = add_para(doc, "3. Biện pháp phát triển năng lực giáo dục hòa nhập của giáo viên phổ thông", "Heading 2", before=8)
measures_heading.paragraph_format.page_break_before = True
add_para(doc, "Bảng 2. Biện pháp phát triển năng lực giáo dục hòa nhập của giáo viên phổ thông", "Table Caption", center=True)
add_pl4_table(doc,
    ["Mục tiêu", "Cách thức thực hiện", "Nguồn lực/ phối hợp/ Điều kiện"],
    [
        ("Xác định đúng điểm xuất phát và một ưu tiên phát triển", "Dùng khung để tự rà soát bằng minh chứng; trao đổi với một đồng nghiệp; chọn tối đa hai chỉ số cần cải thiện trong một học kỳ.", "Thời gian sinh hoạt chuyên môn; mẫu tự rà soát; nguyên tắc bảo mật và không dùng kết quả để xếp hạng."),
        ("Củng cố kiến thức về quyền, rào cản và điều chỉnh hợp lý", "Học mô-đun ngắn; phân tích tình huống; đối chiếu quy định; thực hành chuyển mô tả thiếu hụt thành phân tích rào cản và hỗ trợ.", "Tài liệu chính thống; báo cáo viên/giáo viên cốt cán; tình huống đã ẩn danh; thời gian bồi dưỡng."),
        ("Nâng năng lực thiết kế và đánh giá hòa nhập", "Đồng thiết kế một bài/chủ đề; tạo ít nhất hai cách tham gia hoặc thể hiện; dự giờ tập trung vào việc học; sửa kế hoạch từ minh chứng.", "Tổ/khối chuyên môn; chương trình, SGK; học liệu tiếp cận được; công cụ hỗ trợ; quyền thử nghiệm trong phạm vi được duyệt."),
        ("Tăng chất lượng phối hợp hỗ trợ học sinh", "Thống nhất mục tiêu, vai trò và kênh trao đổi với gia đình, đồng nghiệp, cán bộ hỗ trợ; chỉ chia sẻ thông tin tối thiểu cần thiết; rà tiến độ định kỳ.", "Sự đồng thuận phù hợp; đầu mối hỗ trợ; lịch phối hợp; quy trình lưu trữ và bảo vệ dữ liệu."),
        ("Duy trì cải tiến và lan tỏa thực hành có căn cứ", "Mỗi học kỳ chọn một minh chứng trước–sau; phân tích thay đổi, giới hạn và tải thực hiện; chia sẻ sản phẩm đã ẩn danh; quyết định giữ, sửa hoặc dừng biện pháp.", "Lãnh đạo chuyên môn; đồng nghiệp phản hồi; kho học liệu an toàn; tiêu chí đánh giá khả thi, công bằng và tác động."),
    ],
    [4.0, 6.6, 5.9]
)
add_para(doc, "Nguồn: Tác giả đề xuất.", "Table Source", italic=True)
add_para(doc, (
    "Lưu ý sử dụng: Khung chỉ phục vụ phát triển chuyên môn. Không cộng năm năng lực thành điểm tổng, không công "
    "khai hồ sơ cá nhân và không kết luận năng lực từ một tình huống đơn lẻ. Mọi biện pháp phải được điều chỉnh "
    "theo nhu cầu người học, nguồn lực, thẩm quyền và quy định của đơn vị."
), bold=True, before=4, after=6)

for el in list(body)[start:]:
    refs._p.addprevious(el)

# Replace the old bibliography with sources used by the PL4-aligned report.
for p in list(doc.paragraphs):
    if p._p.getparent() is body and list(body).index(p._p) > list(body).index(refs._p):
        body.remove(p._p)

refs.paragraph_format.page_break_before = True
refs.alignment = WD_ALIGN_PARAGRAPH.CENTER
references = [
    "Bộ Giáo dục và Đào tạo. (2018a). Thông tư số 03/2018/TT-BGDĐT ngày 29 tháng 01 năm 2018 quy định về giáo dục hòa nhập đối với người khuyết tật. https://vanban.chinhphu.vn/?docid=193627&pageid=27160",
    "Bộ Giáo dục và Đào tạo. (2018b). Thông tư số 32/2018/TT-BGDĐT ngày 26 tháng 12 năm 2018 ban hành Chương trình giáo dục phổ thông (được sửa đổi, bổ sung). https://vbpl.vn/bogiaoducdaotao/Pages/vbpq-thuoctinh.aspx?ItemID=146721&Keyword=&dvid=317",
    "Bộ Giáo dục và Đào tạo. (2026). Thông tư số 30/2026/TT-BGDĐT ngày 14 tháng 4 năm 2026 quy định chuẩn nghề nghiệp giáo viên cơ sở giáo dục phổ thông. https://vanban.chinhphu.vn/?classid=0&docid=217839&pageid=27160",
    "Kaplan, I., & Bista, M. B. (2022). Welcoming diversity in the learning environment: Teachers’ handbook for inclusive education. UNESCO. https://unesdoc.unesco.org/ark:/48223/pf0000384009",
    "Miao, F., & Cukurova, M. (2024). AI competency framework for teachers. UNESCO. https://unesdoc.unesco.org/ark:/48223/pf0000391104",
    "Nguyễn Thị Quý. (2024). Xây dựng giờ học hạnh phúc thông qua tương tác tích cực giữa giáo viên và học sinh trong quá trình dạy - học. Vinh University Journal of Science, 53(Special Issue 2), 73–80. https://doi.org/10.56824/vujs.2024.htkhgd75",
    "OECD. (2005). Formative assessment: Improving learning in secondary classrooms. OECD Publishing. https://doi.org/10.1787/9789264007413-en",
    "Quốc hội. (2010). Luật Người khuyết tật số 51/2010/QH12 ngày 17 tháng 6 năm 2010. https://vanban.chinhphu.vn/?docid=96045&pageid=27160",
    "Redecker, C. (2017). European framework for the digital competence of educators: DigCompEdu (Y. Punie, Ed.). Publications Office of the European Union. https://doi.org/10.2760/159770",
    "UNESCO. (2017). A guide for ensuring inclusion and equity in education. https://doi.org/10.54675/MHHZ2237",
]
reference_paragraphs = []
for item in references:
    p = add_para(doc, item, "Normal", after=3)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.first_line_indent = Cm(-0.75)
    reference_paragraphs.append(p)

# In this source package the final section properties make newly appended
# paragraphs land immediately before the retained reference heading. Move the
# completed source list after the heading explicitly and preserve its order.
cursor = refs._p
for p in reference_paragraphs:
    cursor.addnext(p._p)
    cursor = p._p

# Teacher-mandated document geometry and body typography.
for section in doc.sections:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
for style_name in ["Normal", "Figure Caption", "Table Caption", "Table Source", "toc 1", "toc 2"]:
    style = doc.styles[style_name]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
doc.styles["Normal"].paragraph_format.line_spacing = 1.5

doc.core_properties.title = "Khung năng lực giáo dục hòa nhập của giáo viên phổ thông – vận dụng ở tiểu học"
doc.core_properties.subject = "Sản phẩm thực hành theo mẫu PL4"
doc.core_properties.last_modified_by = "Quỳnh Phương"
doc.save(OUTPUT)
print(OUTPUT)
