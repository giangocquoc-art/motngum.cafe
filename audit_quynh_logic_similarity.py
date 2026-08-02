import json
import math
import re
import statistics
import unicodedata
from collections import Counter
from difflib import SequenceMatcher
from pathlib import Path

from docx import Document


ROOT = Path(r"C:\Users\giang\OneDrive\Personal Vault\nghiên cứu khoa học\tiểu luận Quỳnh Phương x3")
TARGET = ROOT / "quynh-phuong-chu-de-4" / "output" / "K51NC_QuynhPhuong_ChuDe4_hoan-chinh.docx"


def extract_docx(path):
    doc = Document(path)
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(c.text.replace("\n", " ").strip() for c in row.cells if c.text.strip())
    return "\n".join(parts)


def norm(text):
    text = unicodedata.normalize("NFC", text.lower())
    text = re.sub(r"https?://\S+", " ", text)
    text = re.sub(r"[^0-9a-zà-ỹđ]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def tokens(text):
    return norm(text).split()


def shingles(words, n=8):
    return {tuple(words[i:i+n]) for i in range(max(0, len(words)-n+1))}


def longest_match(a, b):
    # SequenceMatcher is sufficient for the modest local corpus and reports a
    # contiguous exact-token match after normalization.
    m = SequenceMatcher(None, a, b, autojunk=False).find_longest_match()
    return m.size, " ".join(a[m.a:m.a+m.size])


doc = Document(TARGET)
paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
start = paras.index("PHẦN 1. GIỚI THIỆU VẤN ĐỀ")
end = paras.index("SẢN PHẨM THỰC HÀNH")
report_paras = paras[start:end]
report_prose = [p for p in report_paras if not re.match(r"^(PHẦN\s+\d+|\d+\.\d+\.)", p)]
report_text = " ".join(report_prose)
target_words = tokens(report_text)
target_shingles = shingles(target_words, 8)

# Logic and structure checks.
expected = [
    "PHẦN 1. GIỚI THIỆU VẤN ĐỀ",
    "PHẦN 2. CƠ SỞ LÝ LUẬN",
    "PHẦN 3. THỰC TIỄN VÀ NHU CẦU",
    "PHẦN 4. SẢN PHẨM ĐỀ XUẤT",
    "PHẦN 5. KẾT LUẬN",
]
positions = [paras.index(x) for x in expected]
logic = {
    "section_order_ok": positions == sorted(positions),
    "report_word_count": len(" ".join(report_paras).split()),
    "within_assignment_limit": 1500 <= len(" ".join(report_paras).split()) <= 2500,
    "contains_fabricated_percentage": bool(re.search(r"\b\d+(?:[,.]\d+)?\s*%", report_text)),
    "explicit_scope_limit": "không sử dụng số liệu khảo sát" in report_text.lower(),
    "product_traceability": all(key in report_text for key in ["N1", "N2", "N3", "N4", "N5"]),
}

citations = sorted(set(re.findall(
    r"\b(UNESCO|Kaplan\s+và\s+Bista|OECD|Bộ\s+Giáo\s+dục\s+và\s+Đào\s+tạo)\s*\(((?:19|20)\d{2}[a-z]?)\)",
    report_text,
    flags=re.IGNORECASE,
)))
references_text = "\n".join(paras[paras.index("TÀI LIỆU THAM KHẢO"):])
logic["in_text_citations"] = citations
logic["citations_have_reference_anchor"] = {
    f"{author} ({year})": all(
        name.strip().lower() in references_text.lower()
        for name in re.split(r"\s+và\s+|\s*&\s*", author, flags=re.IGNORECASE)
    ) and year in references_text
    for author, year in citations
}

# Internal repetition scan.
sentences = [s.strip() for s in re.split(r"(?<=[.!?])\s+", report_text) if len(tokens(s)) >= 10]
internal_pairs = []
for i in range(len(sentences)):
    for j in range(i + 1, len(sentences)):
        ratio = SequenceMatcher(None, norm(sentences[i]), norm(sentences[j]), autojunk=False).ratio()
        if ratio >= 0.78:
            internal_pairs.append({"ratio": round(ratio, 3), "a": sentences[i], "b": sentences[j]})

# Local-corpus similarity. Exclude backups, QA outputs and all versions of the
# target author's own directory so the score is not inflated by prior drafts.
candidates = []
for path in ROOT.rglob("*.docx"):
    rel = path.relative_to(ROOT)
    rel_low = str(rel).lower()
    if path == TARGET or "quynh-phuong-chu-de-4" in rel_low or "quynhphuong" in path.name.lower():
        continue
    if any(part in rel_low for part in ["backups", "reports", ".tmp", "_qa"]):
        continue
    candidates.append(path)

comparisons = []
for path in candidates:
    try:
        source_words = tokens(extract_docx(path))
    except Exception:
        continue
    source_shingles = shingles(source_words, 8)
    overlap = target_shingles & source_shingles
    longest_n, excerpt = longest_match(target_words, source_words)
    comparisons.append({
        "file": str(path.relative_to(ROOT)),
        "target_8gram_overlap_pct": round(100 * len(overlap) / max(1, len(target_shingles)), 3),
        "longest_exact_token_run": longest_n,
        "longest_excerpt": excerpt[:300],
    })
comparisons.sort(key=lambda x: (x["target_8gram_overlap_pct"], x["longest_exact_token_run"]), reverse=True)

# AI-style heuristics: this does not infer authorship. It only highlights
# uniform or formulaic prose that merits a human edit.
sentence_lengths = [len(tokens(s)) for s in sentences]
paragraph_lengths = [len(tokens(p)) for p in report_prose if len(tokens(p)) >= 10]
openers = ["trong bối cảnh", "vì vậy", "do đó", "đồng thời", "mặt khác", "thứ nhất", "bên cạnh đó", "tuy nhiên"]
opener_counts = {o: sum(norm(s).startswith(o) for s in sentences) for o in openers}
fourgrams = list(zip(target_words, target_words[1:], target_words[2:], target_words[3:]))
four_counts = Counter(fourgrams)
repeated_fourgram_rate = sum(v - 1 for v in four_counts.values() if v > 1) / max(1, len(fourgrams))
cv = statistics.pstdev(sentence_lengths) / statistics.mean(sentence_lengths) if sentence_lengths else 0
flags = []
if cv < 0.32:
    flags.append("Độ dài câu khá đồng đều; nên đọc lại nhịp văn.")
if max(opener_counts.values(), default=0) >= 4:
    flags.append("Một từ nối đầu câu lặp từ 4 lần trở lên.")
if repeated_fourgram_rate > 0.045:
    flags.append("Tỷ lệ cụm 4 từ lặp lại tương đối cao.")
if not flags:
    flags.append("Không phát hiện dấu hiệu định lượng nổi bật của văn phong quá đồng đều hoặc lặp công thức.")

result = {
    "scope": "Local heuristic audit; not Turnitin and not an AI-authorship detector.",
    "logic": logic,
    "internal_near_duplicate_pairs": internal_pairs,
    "local_similarity_top10": comparisons[:10],
    "ai_style_heuristics": {
        "sentence_count": len(sentence_lengths),
        "sentence_words_mean": round(statistics.mean(sentence_lengths), 2),
        "sentence_words_stdev": round(statistics.pstdev(sentence_lengths), 2),
        "sentence_length_cv": round(cv, 3),
        "paragraph_words_mean": round(statistics.mean(paragraph_lengths), 2),
        "paragraph_words_stdev": round(statistics.pstdev(paragraph_lengths), 2),
        "repeated_4gram_rate": round(repeated_fourgram_rate, 4),
        "transition_opener_counts": opener_counts,
        "flags": flags,
    },
}
print(json.dumps(result, ensure_ascii=False, indent=2))
