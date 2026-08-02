from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


BASE = Path(r"C:\Users\giang\OneDrive\Personal Vault\nghiên cứu khoa học\tiểu luận Quỳnh Phương x3\thu-thao-chu-de-2")
SOURCE = BASE / "output" / "K51NC_ThuThao_ChuDe2.docx"
OUTPUT = BASE / "output" / "K51NC_ThuThao_ChuDe2_ke-hoach-mau-rut-gon.docx"


def set_cell_text(cell, text, bold=False, size=8.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def shade(cell, fill="D9EAD3"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def apply_table_widths(table, widths):
    widths_dxa = [round(width * 567) for width in widths]
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_w = cell._tc.get_or_add_tcPr().tcW
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = Cm(widths[i])
        shade(cell)
        set_cell_text(cell, text, bold=True, size=8.2)
    set_repeat_table_header(table.rows[0])
    for values in rows:
        cells = table.add_row().cells
        for i, text in enumerate(values):
            cells[i].width = Cm(widths[i])
            set_cell_text(cells[i], text, size=8.1)
    apply_table_widths(table, widths)
    return table


def add_para(doc, text="", style="Normal", bold=False, center=False, before=0, after=3):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if text:
        r = p.add_run(text)
        r.bold = bold
    return p


def move_before(paragraph, element):
    paragraph._p.addprevious(element)


def replace_product(doc):
    product = next(p for p in doc.paragraphs if p.text.strip() == "SẢN PHẨM THỰC HÀNH")
    references = next(p for p in doc.paragraphs if p.text.strip() == "TÀI LIỆU THAM KHẢO")
    body = doc._element.body
    children = list(body)
    start = children.index(product._p) + 1
    end = children.index(references._p)
    for element in children[start:end]:
        body.remove(element)

    # New body elements are inserted immediately before the trailing sectPr.
    start_new = len(list(body)) - 1
    # Title block, kept intentionally compact like the group reference plan.
    add_para(doc, "TRƯỜNG TIỂU HỌC CÔNG LẬP ĐÔ THỊ ĐÀ NẴNG (BỐI CẢNH MINH HỌA)", center=True, bold=True, after=1)
    add_para(doc, "TỔ CHUYÊN MÔN KHỐI 4", center=True, bold=True, after=8)
    add_para(doc, "KẾ HOẠCH GIÁO DỤC TÍCH HỢP GIÁO DỤC CÔNG DÂN TOÀN CẦU", center=True, bold=True, after=1)
    add_para(doc, "KHỐI LỚP 4 – NĂM HỌC 2026–2027", center=True, bold=True, after=10)

    add_para(doc, "Tên kế hoạch: “Đà Nẵng xanh – công dân có trách nhiệm”.", bold=True, after=2)
    add_para(doc, "Ý tưởng chung: Học sinh đi qua 5 chặng từ tìm hiểu nơi sống, nhận diện vấn đề nước và môi trường, đối thoại với các góc nhìn, thiết kế hành động nhỏ đến chia sẻ – phản tư. Mỗi chặng được tích hợp vào YCCĐ phù hợp, có sản phẩm học tập và minh chứng ngắn; không tổ chức thành môn học riêng.", after=4)
    add_para(doc, "Sản phẩm học sinh: bản đồ liên hệ địa phương – toàn cầu; phiếu dữ kiện – ý kiến; sơ đồ/bản nháp giải pháp; nhật ký hành động; phần trình bày, phản hồi và bài phản tư. Hình thức có thể thay đổi theo điều kiện lớp, nhưng không thay đổi YCCĐ, nguyên tắc an toàn và bảo vệ dữ liệu.", after=8)

    add_para(doc, "I. ĐẶC ĐIỂM TÌNH HÌNH", style="Heading 3", after=4)
    add_para(doc, "1. Khái quát tình hình chung", style="Heading 4", after=2)
    add_para(doc, "Kế hoạch dành cho khối 4 tại một trường tiểu học công lập đô thị ở Đà Nẵng. GCED được tích hợp vào các môn học và hoạt động giáo dục hiện có, lấy vấn đề gần gũi về nơi sống, nước, môi trường, văn hóa và thông tin số làm điểm xuất phát.", after=3)
    add_para(doc, "2. Thuận lợi và khó khăn", style="Heading 4", after=2)
    for text in [
        "Thuận lợi: Chương trình theo định hướng phát triển phẩm chất, năng lực; học sinh hứng thú với hoạt động tìm hiểu, hợp tác, vẽ/trình bày và hành động vừa sức.",
        "Khó khăn: Điều kiện thiết bị, học liệu và thời gian phối hợp giữa các lớp có thể khác nhau; dễ phát sinh hoạt động hình thức nếu không bám YCCĐ và minh chứng.",
        "Yêu cầu trước khi áp dụng: rà đặc điểm người học, bộ sách và tiến độ môn học, học liệu – thiết bị, an toàn/bảo vệ trẻ em, địa điểm và đầu mối phối hợp. Thiếu điều kiện thì thu hẹp hoặc thay bằng phương án tại lớp.",
    ]:
        add_para(doc, text, style="List Bullet", after=1)

    add_para(doc, "II. MỤC TIÊU", style="Heading 3", before=5, after=4)
    add_para(doc, "Giúp học sinh nhận biết liên hệ giữa bản thân, Đà Nẵng và thế giới; biết tìm hiểu thông tin, lắng nghe khác biệt, hợp tác và lựa chọn hành động có trách nhiệm. Kế hoạch góp phần phát triển phẩm chất nhân ái, trách nhiệm; năng lực giao tiếp – hợp tác, giải quyết vấn đề, tự học và năng lực đặc thù của môn học.", after=3)
    for text in [
        "Nhận thức: giải thích bằng ví dụ vừa sức liên hệ địa phương – toàn cầu; phân biệt dữ kiện với ý kiến ở ngữ liệu được cung cấp.",
        "Cảm xúc – xã hội: lắng nghe, tôn trọng khác biệt và giới thiệu bản sắc địa phương không khuôn mẫu hóa cộng đồng khác.",
        "Hành vi: cùng nhóm lập kế hoạch, thực hiện và phản tư về một hành động nhỏ, an toàn, có căn cứ.",
        "Công dân số: kiểm tra tác giả, thời điểm và bằng chứng; bảo vệ dữ liệu cá nhân; giao tiếp lịch sự, ghi nguồn khi sử dụng nội dung/hình ảnh.",
    ]:
        add_para(doc, text, style="List Bullet", after=1)

    add_para(doc, "III. KẾ HOẠCH THỰC HIỆN", style="Heading 3", before=5, after=4)
    add_para(doc, "Kế hoạch được chuẩn bị trong tháng 5–8/2026: rà YCCĐ, khóa điểm tích hợp theo bộ sách và lịch trường, chuẩn bị học liệu – công cụ đánh giá, phê duyệt an toàn. Từ tháng 9/2026 đến tháng 5/2027, lớp triển khai 5 chặng dưới đây. Sau mỗi chặng, tổ khối quyết định giữ, sửa hoặc dừng dựa trên YCCĐ, mức độ tham gia, an toàn và minh chứng thực tế.", after=4)
    add_para(doc, "Bảng 1. Tóm tắt hành trình 5 chặng", style="Table Caption", after=2)
    table1 = add_table(doc,
        ["Chặng", "Thời gian", "Trọng tâm", "Minh chứng chính"],
        [
            ("1. Khởi động", "Tháng 9", "Nơi sống, bản sắc, quy tắc đối thoại; khảo sát điểm xuất phát.", "Bản đồ ý tưởng; quy tắc lớp; phiếu nhu cầu."),
            ("2. Kết nối", "Tháng 10–11", "Đà Nẵng – Việt Nam – thế giới; tôn trọng đa dạng văn hóa và thông tin.", "Bản đồ liên hệ; phiếu dữ kiện – ý kiến; phần giới thiệu."),
            ("3. Hiểu vấn đề", "Tháng 12–1", "Nước là tài nguyên chung; nguyên nhân – hệ quả của ô nhiễm và sử dụng lãng phí.", "Sơ đồ nước; bảng nguyên nhân – hệ quả; so sánh giải pháp."),
            ("4. Hành động", "Tháng 2–4", "Đồng thiết kế, thử nghiệm hành động xanh vừa sức; phản hồi và điều chỉnh.", "Kế hoạch nhóm; phân vai; nhật ký hành động; phản hồi."),
            ("5. Lan tỏa", "Tháng 5", "Chia sẻ kết quả, tự đánh giá, phản tư và đề xuất năm sau.", "Sản phẩm cuối; bài trình bày; bài phản tư; danh mục giữ – sửa – dừng."),
        ], [2.3, 2.0, 7.3, 4.9])
    add_para(doc, "Bảng 2. Kế hoạch tích hợp theo các chặng", style="Table Caption", before=5, after=2)
    table2 = add_table(doc,
        ["Nội dung tích hợp", "Mục tiêu/YCCĐ", "Môn tích hợp", "Địa chỉ/mạch học", "PP, HTTC", "Điều kiện thực hiện"],
        [
            ("Địa phương – toàn cầu; bản sắc và trách nhiệm nơi sống", "Xác định vị trí, mô tả nét chính của địa phương; thể hiện tình cảm và trách nhiệm với môi trường xung quanh.", "Lịch sử và Địa lí; HĐTN", "Địa phương em; hoạt động vì cộng đồng", "Đọc bản đồ/tư liệu; nghĩ – trao đổi – chia sẻ", "Bản đồ, tư liệu chính thức; bản in khi thiếu thiết bị."),
            ("Nước trong đời sống và sản xuất", "Nêu vai trò, tính chất của nước; mô tả vòng tuần hoàn; liên hệ thực tế phù hợp.", "Khoa học", "Chủ đề Nước", "Thí nghiệm đơn giản; sơ đồ hóa; thảo luận", "Dụng cụ an toàn; không yêu cầu dữ liệu riêng của gia đình."),
            ("Bảo vệ nguồn nước và tiêu dùng có trách nhiệm", "Nêu nguyên nhân ô nhiễm; đề xuất cách bảo vệ, sử dụng nước tiết kiệm; giải thích hệ quả.", "Khoa học; Đạo đức", "Chủ đề Nước; bảo vệ của công/trách nhiệm", "Nghiên cứu trường hợp; so sánh giải pháp", "Chỉ dùng mẫu/hình ảnh an toàn; hành động được phê duyệt."),
            ("Hành động xanh và hợp tác", "Lập kế hoạch nhỏ, phân vai, thực hiện an toàn; tiếp nhận phản hồi và điều chỉnh.", "HĐTN; Mĩ thuật", "Tìm hiểu và bảo vệ môi trường; tạo hình/truyền thông", "Dự án nhỏ; thiết kế poster/sơ đồ; phản tư", "Vật liệu sạch, tái sử dụng; có phương án không dùng thiết bị."),
            ("Công dân số và chia sẻ có trách nhiệm", "Kiểm tra nguồn cơ bản; ghi nguồn; bảo vệ dữ liệu; giao tiếp tôn trọng khi trình bày/chia sẻ.", "Tích hợp trong các chặng", "Hoạt động tìm hiểu, trình bày và phản hồi", "Phiếu kiểm tra nguồn; phản hồi đồng đẳng", "Không đăng công khai nếu chưa được duyệt; có lựa chọn tham gia tương đương."),
        ], [2.7, 3.4, 2.0, 2.8, 2.4, 3.2])
    add_para(doc, "Lưu ý: Tổ chuyên môn bổ sung tên bài, tuần và tiết sau khi trường khóa bộ sách và kế hoạch năm học; không tự tạo thêm môn học hoặc yêu cầu ngoài YCCĐ.", after=5)

    add_para(doc, "IV. TỔ CHỨC THỰC HIỆN", style="Heading 3", before=4, after=4)
    for text in [
        "Ban giám hiệu/lãnh đạo chuyên môn: phê duyệt mục tiêu, tiến độ, điều kiện an toàn và phối hợp; không triển khai hoạt động khi điều kiện tối thiểu chưa bảo đảm.",
        "Tổ chuyên môn khối 4: rà YCCĐ, thống nhất điểm tích hợp, điều phối học liệu và họp định kỳ để xem minh chứng, quyết định giữ – sửa – dừng.",
        "Giáo viên: thiết kế nhiệm vụ vừa sức, phân hóa hỗ trợ, hướng dẫn ghi nhận minh chứng và phản hồi quá trình; không biến sản phẩm thành áp lực thi đua.",
        "Học sinh: lựa chọn vai trò, hợp tác, thực hiện nhiệm vụ an toàn, tự đánh giá và phản tư; không tự liên hệ với tổ chức bên ngoài.",
        "Cha mẹ/cộng đồng: phối hợp tự nguyện, đúng phạm vi trường duyệt; không phải cung cấp dữ liệu, hình ảnh hoặc đóng góp vật chất để học sinh được tham gia.",
    ]:
        add_para(doc, text, style="List Bullet", after=1)
    add_para(doc, "Bảo đảm an toàn và hòa nhập: có bản giấy/hình ảnh dễ đọc và phương án tham gia thay thế; không thu dữ liệu định danh không cần thiết; không xử lý rác nguy hại hoặc tổ chức hoạt động gần nguồn nước/ngoài trường nếu chưa được phê duyệt, giám sát.", after=4)

    add_para(doc, "V. ĐÁNH GIÁ", style="Heading 3", before=4, after=4)
    add_para(doc, "Đánh giá vì sự tiến bộ, bám YCCĐ và sử dụng minh chứng trong suốt quá trình; kết hợp nhận xét của giáo viên, tự đánh giá và đánh giá đồng đẳng. Không chấm thái độ chính trị, hoàn cảnh gia đình hoặc mức độ đồng thuận với giáo viên.", after=3)
    for text in [
        "Thời điểm: đầu vào (bản đồ ý tưởng/câu hỏi tình huống); trong quá trình (quan sát, phiếu học tập, bản nháp, phản hồi); cuối chu trình (hồ sơ học tập, sản phẩm, trình bày, bài phản tư).",
        "Nội dung: hiểu biết và kiểm tra nguồn; đồng cảm và đối thoại; hành động an toàn, có căn cứ; hợp tác và công dân số.",
        "Minh chứng: bản đồ liên hệ, phiếu dữ kiện – ý kiến, sơ đồ/bảng phân tích, kế hoạch và nhật ký hành động, sản phẩm truyền thông có ghi nguồn, phản hồi và bài phản tư.",
        "Cuối mỗi chặng, tổ khối xem mức độ bám YCCĐ, cơ hội tham gia, an toàn, chất lượng minh chứng và tải công việc để điều chỉnh; chỉ báo cáo hiệu quả khi có bằng chứng thực tế.",
    ]:
        add_para(doc, text, style="List Bullet", after=1)

    # Move all newly appended elements from after the reference list to the product location.
    new_elements = list(body)[start_new:]
    for element in new_elements:
        move_before(references, element)


def update_front_matter(doc):
    # Replace the obsolete manual list of tables with the two compact plan tables.
    for p in doc.paragraphs:
        t = " ".join(p.text.split())
        if t.startswith("Bảng 1. Ma trận tích hợp giáo dục công dân toàn cầu"):
            p.text = "Bảng 1. Tóm tắt hành trình 5 chặng 9"
        elif t.startswith("Bảng 2. Tiến trình xây dựng và triển khai"):
            p.text = "Bảng 2. Kế hoạch tích hợp theo các chặng 9"
        elif t.startswith("Bảng 3. Mô tả bốn mức đánh giá"):
            p._element.getparent().remove(p._element)
        elif p.style.name == "toc 1" and t.startswith("SẢN PHẨM THỰC HÀNH"):
            p.text = "SẢN PHẨM THỰC HÀNH\t8"
        elif p.style.name == "toc 1" and t.startswith("TÀI LIỆU THAM KHẢO"):
            p.text = "TÀI LIỆU THAM KHẢO\t10"


doc = Document(SOURCE)
replace_product(doc)
update_front_matter(doc)
doc.save(OUTPUT)
print(OUTPUT)
