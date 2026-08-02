from pathlib import Path
from copy import deepcopy

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

BASE = Path(r"C:\Users\giang\OneDrive\Personal Vault\nghiên cứu khoa học\tiểu luận Quỳnh Phương x3\thu-thao-chu-de-2")
SOURCE = BASE / "output" / "K51NC_ThuThao_ChuDe2.docx"
OUTPUT = BASE / "output" / "K51NC_ThuThao_ChuDe2_trong-tam-san-pham.docx"


def para(doc, text="", style="Normal", center=False, bold=False, after=3, before=0):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    if text:
        r = p.add_run(text)
        r.bold = bold
    return p


def shade(cell):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "E7E6E6")
    cell._tc.get_or_add_tcPr().append(shd)


def repeat_header(row):
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(marker)


def set_widths(table, widths):
    widths = [round(x * 567) for x in widths]
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tcw = cell._tc.get_or_add_tcPr().tcW
            tcw.set(qn("w:w"), str(widths[index]))
            tcw.set(qn("w:type"), "dxa")


def cell_text(cell, text, bold=False, size=10.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    # Narrow PL2 columns become difficult to read when Word inherits the
    # document-wide justified alignment (it stretches or breaks Vietnamese
    # words).  Keep table body text left-aligned; header cells are centered
    # explicitly below.
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(doc, headers, rows, widths, size=10.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, header in enumerate(headers):
        shade(t.rows[0].cells[i])
        cell_text(t.rows[0].cells[i], header, True, size)
        t.rows[0].cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    repeat_header(t.rows[0])
    for values in rows:
        cells = t.add_row().cells
        for i, value in enumerate(values):
            cell_text(cells[i], value, size=size)
    set_widths(t, widths)
    return t


def add_figure(doc, image_path, caption, width_cm, source_text):
    para(doc, caption, "Figure Caption", after=2)
    picture = doc.add_picture(str(image_path), width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(doc, source_text, "Normal", center=True, after=5)


def remove_table_borders(t):
    tbl_pr = t._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)


def set_admin_cell(cell, lines, bold_lines=(), italic_lines=()):
    cell.text = ""
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(line)
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)
        r.bold = i in bold_lines
        r.italic = i in italic_lines


def add_admin_header(doc):
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    set_admin_cell(t.cell(0, 0), ["TRƯỜNG TIỂU HỌC …………………", "TỔ CHUYÊN MÔN KHỐI 4"], bold_lines=(0, 1))
    set_admin_cell(t.cell(0, 1), ["CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", "Độc lập – Tự do – Hạnh phúc", "____________________", "Đà Nẵng, ngày ….. tháng ….. năm 2026"], bold_lines=(0, 1), italic_lines=(3,))
    set_widths(t, [7.8, 8.7])
    remove_table_borders(t)
    para(doc, "", after=4)


def add_signature_block(doc):
    para(doc, "", after=5)
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    set_admin_cell(t.cell(0, 0), ["TỔ TRƯỞNG CHUYÊN MÔN", "(Ký, ghi rõ họ tên)"], bold_lines=(0,), italic_lines=(1,))
    set_admin_cell(t.cell(0, 1), ["NGƯỜI LẬP KẾ HOẠCH", "(Ký, ghi rõ họ tên)"], bold_lines=(0,), italic_lines=(1,))
    set_widths(t, [8.25, 8.25])
    remove_table_borders(t)
    for _ in range(3):
        para(doc, "", after=0)


def insert_product(doc):
    product = next(p for p in doc.paragraphs if p.text.strip() == "SẢN PHẨM THỰC HÀNH")
    refs = next(p for p in doc.paragraphs if p.text.strip() == "TÀI LIỆU THAM KHẢO")
    body = doc._element.body
    children = list(body)
    for el in children[children.index(product._p) + 1: children.index(refs._p)]:
        body.remove(el)
    start = len(list(body)) - 1

    add_admin_header(doc)
    para(doc, "KẾ HOẠCH GIÁO DỤC TÍCH HỢP GIÁO DỤC CÔNG DÂN TOÀN CẦU", center=True, bold=True, after=1)
    para(doc, "KHỐI LỚP 4 – NĂM HỌC 2026–2027", center=True, bold=True, after=9)
    para(doc, "Tên kế hoạch: “Dòng chảy xanh – Bản đồ nước và trách nhiệm”.", bold=True, after=2)
    para(doc, "Ý tưởng chung: Mỗi lớp cùng tạo một “Bản đồ dòng chảy xanh” mô tả hành trình của nước từ thiên nhiên đến gia đình, trường học và trở lại môi trường. Trên bản đồ, học sinh gắn các điểm quan sát, câu hỏi, góc nhìn và hành động nhỏ. Bản đồ được hoàn thiện dần qua bốn trạm học tập, không phải sản phẩm làm một lần để trưng bày.", after=3)
    para(doc, "Sản phẩm chung gồm: bản đồ dòng chảy xanh của lớp; thẻ “Một giọt nước kể chuyện” ghi dữ kiện/ý kiến; bảng nguyên nhân – hệ quả; kế hoạch hành động nhỏ; góc chia sẻ có ghi nguồn và bài phản tư. Sản phẩm được xây dựng riêng cho bối cảnh khối 4 và được cập nhật trong suốt năm học.", after=3)
    para(doc, "Bộ sản phẩm nộp kèm kế hoạch gồm:", bold=True, after=2)
    for text in [
        "01 bản kế hoạch giáo dục tích hợp theo cấu trúc PL2.",
        "01 mẫu nền Bản đồ dòng chảy xanh khổ A1 cho mỗi lớp, gồm tuyến nước và bốn khu vực cập nhật minh chứng.",
        "Bộ thẻ học tập: “Một giọt nước kể chuyện”, dữ kiện – ý kiến, nguyên nhân – hệ quả, kế hoạch hành động và phản tư.",
        "Bảng kiểm ngắn cho giáo viên và học sinh, dùng để theo dõi YCCĐ, mức độ tham gia, an toàn và việc ghi nguồn.",
    ]:
        para(doc, text, "List Bullet", after=1)
    add_figure(doc, BASE / "figures" / "auto" / "dong-chay-xanh-ban-do-mau.png",
               "Hình 2. Mẫu Bản đồ dòng chảy xanh của lớp 4", 14.5,
               "Nguồn: Hình minh họa AI do tác giả tạo cho sản phẩm đề xuất.")
    add_figure(doc, BASE / "figures" / "auto" / "dong-chay-xanh-bo-the-mau.png",
               "Hình 3. Mẫu bộ thẻ học tập đi kèm bản đồ", 8.8,
               "Nguồn: Hình minh họa AI do tác giả tạo cho sản phẩm đề xuất.")

    para(doc, "I. ĐẶC ĐIỂM TÌNH HÌNH", "Heading 3", after=3)
    para(doc, "1. Khái quát tình hình chung", "Heading 4", after=2)
    para(doc, "Kế hoạch dành cho khối 4 tại trường tiểu học công lập đô thị ở Đà Nẵng. Giáo dục công dân toàn cầu được tích hợp vào môn học và hoạt động giáo dục hiện có, lấy nước – một tài nguyên chung, gần gũi với đời sống địa phương – làm tình huống học tập.", after=3)
    para(doc, "2. Thuận lợi và khó khăn", "Heading 4", after=2)
    for text in [
        "Thuận lợi: học sinh có thể quan sát, thí nghiệm, kể chuyện, vẽ sơ đồ và hợp tác để tìm hiểu nước trong đời sống hằng ngày.",
        "Khó khăn: điều kiện thiết bị, thời gian phối hợp và khả năng tham gia giữa các lớp không giống nhau; hoạt động dễ thành phong trào nếu chỉ chú ý đến áp phích đẹp.",
        "Điều kiện áp dụng: tổ khối rà bộ sách, tiến độ môn học, nhu cầu hỗ trợ, học liệu – thiết bị, an toàn, địa điểm và đầu mối phối hợp; thiếu điều kiện thì dùng ngữ liệu/hình ảnh an toàn tại lớp.",
    ]:
        para(doc, text, "List Bullet", after=1)

    para(doc, "II. MỤC TIÊU", "Heading 3", before=4, after=3)
    para(doc, "Giúp học sinh nhận biết nước gắn kết con người, địa phương và thế giới; biết tìm hiểu thông tin, lắng nghe khác biệt, hợp tác và thực hiện hành động nhỏ có trách nhiệm. Kế hoạch hỗ trợ YCCĐ môn học, góp phần phát triển phẩm chất trách nhiệm, nhân ái và năng lực giao tiếp – hợp tác, giải quyết vấn đề, tự học.", after=3)
    for text in [
        "Nhận thức: giải thích được vai trò, vòng tuần hoàn và một số vấn đề sử dụng/bảo vệ nước; phân biệt dữ kiện với ý kiến trong ngữ liệu.",
        "Cảm xúc – xã hội: lắng nghe nhu cầu, góc nhìn khác nhau về nước; tôn trọng sự khác biệt và không đổ lỗi cho cá nhân/hoàn cảnh.",
        "Hành vi: đề xuất, thử nghiệm và phản tư về một hành động tiết kiệm hoặc bảo vệ nước an toàn, vừa sức.",
        "Công dân số: kiểm tra nguồn, ghi nguồn cho hình ảnh/nội dung, bảo vệ dữ liệu cá nhân và chia sẻ lịch sự.",
    ]:
        para(doc, text, "List Bullet", after=1)

    para(doc, "III. KẾ HOẠCH THỰC HIỆN", "Heading 3", before=4, after=3)
    para(doc, "Tháng 5–8/2026, tổ khối rà YCCĐ, khóa địa chỉ tích hợp theo bộ sách, chuẩn bị học liệu, công cụ đánh giá và phương án an toàn. Từ tháng 9/2026 đến tháng 5/2027, mỗi lớp triển khai bốn trạm dưới đây. Sau mỗi trạm, giáo viên và tổ khối xem minh chứng để giữ, sửa hoặc dừng hoạt động.", after=4)
    para(doc, "Bảng 1. Hành trình bốn trạm của “Dòng chảy xanh”", "Table Caption", after=2)
    table(doc,
        ["Trạm", "Thời gian", "Việc học của học sinh", "Dấu vết học tập"],
        [
            ("1. Theo dấu một giọt nước", "Tháng 9–10", "Quan sát nước quanh em; dựng phần đầu của bản đồ; nêu câu hỏi và phân biệt dữ kiện – ý kiến.", "Bản đồ ý tưởng; thẻ “Một giọt nước kể chuyện”; quy tắc đối thoại."),
            ("2. Lắng nghe dòng chảy", "Tháng 11–12", "Tìm hiểu nước ở Đà Nẵng, trong sản xuất và đời sống; so sánh các góc nhìn, nguyên nhân – hệ quả.", "Bản đồ dòng chảy; bảng nguyên nhân – hệ quả; phiếu kiểm tra nguồn."),
            ("3. Thử một thay đổi nhỏ", "Tháng 1–3", "Nhóm chọn một việc vừa sức: dùng nước hợp lí, bảo vệ nguồn nước hoặc truyền thông tại lớp; ghi nhận và điều chỉnh.", "Kế hoạch nhóm; phân vai; nhật ký hành động; phản hồi đồng đẳng."),
            ("4. Gửi dòng chảy đi xa", "Tháng 4–5", "Hoàn thiện bản đồ, chia sẻ điều đã học và điều chưa làm được; tự đánh giá, phản tư.", "Góc chia sẻ có ghi nguồn; sản phẩm cuối; bài phản tư; đề xuất năm sau."),
        ], [2.8, 2.2, 6.3, 5.2])
    para(doc, "Bảng 2. Kế hoạch tích hợp theo các trạm", "Table Caption", before=5, after=2)
    table(doc,
        ["Nội dung", "Mục tiêu/YCCĐ", "Môn tích hợp", "Tên bài/chủ đề tích hợp", "PP, HTTC", "Điều kiện"],
        [
            ("Nước và nơi sống; liên hệ địa phương – toàn cầu", "Nêu vai trò, tính chất của nước; mô tả vòng tuần hoàn; liên hệ Đà Nẵng với những nơi khác qua cùng một tài nguyên.", "Khoa học; Lịch sử và Địa lí", "Chủ đề Nước; Địa phương em", "Quan sát, thí nghiệm đơn giản, sơ đồ hóa", "Dụng cụ an toàn; bản đồ/tư liệu chính thức; bản in khi thiếu thiết bị."),
            ("Nước là tài nguyên chung; nhiều góc nhìn", "Nêu nguyên nhân ô nhiễm, giải thích hệ quả; so sánh nhu cầu/góc nhìn của gia đình, trường học và cộng đồng; đề xuất cách bảo vệ, sử dụng tiết kiệm.", "Khoa học; Đạo đức", "Bảo vệ nước; trách nhiệm với của công/môi trường", "Nghiên cứu tình huống; đối thoại; so sánh giải pháp", "Chỉ dùng mẫu/hình ảnh an toàn; không tiếp xúc nguồn nước ô nhiễm."),
            ("Bản đồ hành động xanh", "Lập kế hoạch nhỏ, phân vai, thực hiện an toàn; tiếp nhận phản hồi để điều chỉnh.", "HĐTN; Mĩ thuật", "Tìm hiểu và bảo vệ môi trường; tạo hình/truyền thông", "Dự án nhỏ; thiết kế bản đồ/góc chia sẻ", "Vật liệu sạch, tái sử dụng; phương án không dùng thiết bị."),
            ("Thông tin và chia sẻ có trách nhiệm", "Kiểm tra tác giả, thời điểm, bằng chứng; ghi nguồn; bảo vệ dữ liệu và giao tiếp tôn trọng.", "Tích hợp trong các trạm", "Tìm hiểu, trình bày và phản hồi", "Phiếu kiểm tra nguồn; phản hồi đồng đẳng", "Không công bố hình ảnh/dữ liệu khi chưa được duyệt; có lựa chọn tham gia tương đương."),
        ], [2.6, 2.9, 2.5, 2.65, 2.55, 3.3], size=9.5)
    para(doc, "Lưu ý: tên bài, tuần và tiết sẽ được tổ chuyên môn bổ sung sau khi trường chốt bộ sách và kế hoạch năm học; không tạo thêm môn học hoặc yêu cầu ngoài YCCĐ.", after=4)

    para(doc, "IV. TỔ CHỨC THỰC HIỆN", "Heading 3", before=4, after=3)
    for text in [
        "Ban giám hiệu/lãnh đạo chuyên môn: phê duyệt tiến độ, điều kiện an toàn, học liệu và phối hợp; không duyệt hoạt động khi điều kiện tối thiểu chưa bảo đảm.",
        "Tổ chuyên môn khối 4: thống nhất YCCĐ và công cụ; điều phối học liệu; xem minh chứng định kì để quyết định giữ – sửa – dừng.",
        "Giáo viên: thiết kế nhiệm vụ vừa sức, phân hóa hỗ trợ, hướng dẫn học sinh cập nhật bản đồ và phản hồi quá trình; không dùng bản đồ để xếp hạng lớp/học sinh.",
        "Học sinh: chọn vai trò, hợp tác, thực hiện nhiệm vụ an toàn, ghi nhận điều học được và tự phản tư.",
        "Cha mẹ/cộng đồng: phối hợp tự nguyện, đúng phạm vi được trường duyệt; không phải cung cấp số liệu sử dụng nước, hình ảnh hoặc đóng góp vật chất để học sinh được tham gia.",
    ]:
        para(doc, text, "List Bullet", after=1)
    para(doc, "Bảo đảm an toàn và hòa nhập: có bản giấy, hình ảnh dễ đọc và vai trò thay thế; không thu dữ liệu định danh không cần thiết; không xử lý rác nguy hại, di chuyển ngoài trường hoặc hoạt động gần nguồn nước khi chưa được phê duyệt, giám sát.", after=4)

    para(doc, "V. ĐÁNH GIÁ", "Heading 3", before=4, after=3)
    para(doc, "Đánh giá vì sự tiến bộ, bám YCCĐ và dựa trên minh chứng trong cả bốn trạm; kết hợp nhận xét của giáo viên, tự đánh giá và đánh giá đồng đẳng. Không chấm hoàn cảnh gia đình, mức tiêu thụ nước của gia đình hay sự đồng thuận với giáo viên.", after=3)
    for text in [
        "Thời điểm: đầu vào (câu hỏi, bản đồ ý tưởng); trong quá trình (quan sát, thẻ dữ kiện – ý kiến, bản nháp, nhật ký); cuối chu trình (bản đồ hoàn thiện, phần chia sẻ, bài phản tư).",
        "Nội dung: hiểu biết và kiểm tra nguồn; đồng cảm và đối thoại; hành động an toàn, có căn cứ; hợp tác và công dân số.",
        "Minh chứng: bản đồ dòng chảy xanh, thẻ “Một giọt nước kể chuyện”, bảng nguyên nhân – hệ quả, kế hoạch/nhật ký hành động, sản phẩm có ghi nguồn, phản hồi và bài phản tư.",
        "Cuối mỗi trạm, tổ khối xem mức độ bám YCCĐ, cơ hội tham gia, an toàn, chất lượng minh chứng và tải công việc để điều chỉnh; chỉ báo cáo hiệu quả khi có bằng chứng thực tế.",
    ]:
        para(doc, text, "List Bullet", after=1)
    add_signature_block(doc)

    # A formal PL2 plan ends at the signature block.  Start the academic
    # reference list on its own page and avoid stretched spacing in URLs and
    # long Vietnamese titles.
    refs.paragraph_format.page_break_before = True
    refs.alignment = WD_ALIGN_PARAGRAPH.CENTER
    refs_index = next(i for i, p in enumerate(doc.paragraphs) if p._p is refs._p)
    for p in doc.paragraphs[refs_index + 1:]:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for el in list(body)[start:]:
        refs._p.addprevious(el)


def update_front(doc, table1_page=None, table2_page=None, refs_page=None, figure2_page=None, figure3_page=None):
    for p in doc.paragraphs:
        text = " ".join(p.text.split())
        if text.startswith("Bảng 1. Ma trận tích hợp giáo dục công dân toàn cầu"):
            p.text = "Bảng 1. Hành trình bốn trạm của Dòng chảy xanh" + (f"\t{table1_page}" if table1_page else "")
        elif text.startswith("Bảng 2. Tiến trình xây dựng và triển khai"):
            p.text = "Bảng 2. Kế hoạch tích hợp theo các trạm" + (f"\t{table2_page}" if table2_page else "")
        elif text.startswith("Bảng 3. Mô tả bốn mức đánh giá"):
            p._element.getparent().remove(p._element)
        elif p.style.name == "toc 1" and text.startswith("SẢN PHẨM THỰC HÀNH"):
            p.text = "SẢN PHẨM THỰC HÀNH\t8"
        elif p.style.name == "toc 1" and text.startswith("TÀI LIỆU THAM KHẢO") and refs_page:
            p.text = f"TÀI LIỆU THAM KHẢO\t{refs_page}"
        elif p.style.name == "Normal" and text.startswith("Hình 2. Mẫu Bản đồ dòng chảy xanh"):
            p.text = "Hình 2. Mẫu Bản đồ dòng chảy xanh của lớp 4" + (f"\t{figure2_page}" if figure2_page else "")
        elif p.style.name == "Normal" and text.startswith("Hình 3. Mẫu bộ thẻ học tập"):
            p.text = "Hình 3. Mẫu bộ thẻ học tập đi kèm bản đồ" + (f"\t{figure3_page}" if figure3_page else "")

    # The source file only had Figure 1 in its manually prepared list.  Clone
    # that list-entry paragraph so its dotted tab leader and tab stop are kept.
    figure_list_entry = next(
        p for p in doc.paragraphs
        if p.style.name == "Normal" and p.text.strip().startswith("Hình 1. Tiến trình phát triển sản phẩm")
    )
    h2_el = deepcopy(figure_list_entry._p)
    figure_list_entry._p.addnext(h2_el)
    h2 = Paragraph(h2_el, figure_list_entry._parent)
    h2.text = "Hình 2. Mẫu Bản đồ dòng chảy xanh của lớp 4" + (f"\t{figure2_page}" if figure2_page else "")
    h3_el = deepcopy(figure_list_entry._p)
    h2._p.addnext(h3_el)
    h3 = Paragraph(h3_el, figure_list_entry._parent)
    h3.text = "Hình 3. Mẫu bộ thẻ học tập đi kèm bản đồ" + (f"\t{figure3_page}" if figure3_page else "")


def align_report_with_product(doc):
    replacements = {
        "XÂY DỰNG KẾ HOẠCH GIÁO DỤC KHỐI 4 TÍCH HỢP": "XÂY DỰNG KẾ HOẠCH GIÁO DỤC KHỐI 4 TÍCH HỢP GIÁO DỤC CÔNG DÂN TOÀN CẦU QUA SẢN PHẨM ‘DÒNG CHẢY XANH – BẢN ĐỒ NƯỚC VÀ TRÁCH NHIỆM’",
        "Giáo dục công dân toàn cầu (Global Citizenship Education": "Giáo dục công dân toàn cầu (Global Citizenship Education — GCED) dễ bị thu hẹp thành khẩu hiệu hoặc ngày hội có sản phẩm đẹp nhưng không cho biết học sinh đã hiểu, đối thoại và hành động ra sao. Bài viết vì vậy tập trung xây dựng một sản phẩm cụ thể cho khối 4: “Dòng chảy xanh – Bản đồ nước và trách nhiệm”. Sản phẩm chỉ có giá trị khi nối rõ YCCĐ – nhiệm vụ của học sinh – minh chứng – đánh giá.",
        "Sản phẩm được thiết kế theo mẫu PL2": "Sản phẩm được thiết kế theo cấu trúc PL2 cho một trường tiểu học công lập đô thị tại Đà Nẵng, năm học 2026–2027, ở dạng minh họa. Phần lý luận và thực tiễn chỉ giữ vai trò giải thích các quyết định thiết kế; trọng tâm của bài là cấu tạo sản phẩm, tiến trình bốn trạm, địa chỉ tích hợp, cách tổ chức và minh chứng đánh giá. Do chưa có dữ liệu trường cụ thể, bài không tự gán quy mô, nguồn lực hoặc kết quả.",
        "Vấn đề chính không phải thiếu hoạt động môi trường": "Vấn đề chính không phải thiếu hoạt động môi trường, mà là các sản phẩm thường rời rạc, thiên về trang trí và thiếu dấu vết cho thấy học sinh đã học gì. “Bản đồ dòng chảy xanh” được chọn để gom minh chứng của nhiều môn vào một sản phẩm phát triển dần: học sinh bổ sung dữ kiện, góc nhìn, nguyên nhân – hệ quả, hành động và phản tư theo từng trạm. Nhờ đó, giáo viên có thể nhìn thấy tiến bộ thay vì chỉ chấm áp phích cuối cùng.",
        "Sản phẩm là Kế hoạch giáo dục khối 4": "Sản phẩm đề xuất gồm hai lớp gắn với nhau: (1) Kế hoạch giáo dục khối 4 tích hợp GCED theo PL2; (2) bộ sản phẩm học tập “Dòng chảy xanh – Bản đồ nước và trách nhiệm”. Bản đồ không phải tranh trang trí mà là hồ sơ học tập trực quan, được cập nhật qua bốn trạm: theo dấu một giọt nước; lắng nghe dòng chảy; thử một thay đổi nhỏ; gửi dòng chảy đi xa.",
        "Ma trận chọn sáu địa chỉ": "Ma trận tập trung bốn cụm nội dung: nước và nơi sống; nước là tài nguyên chung; bản đồ hành động xanh; thông tin và chia sẻ có trách nhiệm. Các cụm nối YCCĐ của Khoa học, Lịch sử và Địa lí, Đạo đức, Hoạt động trải nghiệm và Mĩ thuật với nhiệm vụ cụ thể trên bản đồ. Tên bài, tuần và tiết được bổ sung sau khi trường khóa bộ sách và kế hoạch môn học.",
        "Kế hoạch bắt đầu bằng xây dựng dự thảo tháng 5–6": "Kế hoạch được chuẩn bị trong tháng 5–8 bằng việc rà YCCĐ, chốt điều kiện, học liệu và công cụ. Trong năm học, Trạm 1 diễn ra tháng 9–10; Trạm 2 tháng 11–12; Trạm 3 tháng 1–3; Trạm 4 tháng 4–5. Mỗi trạm chỉ hoàn tất khi bản đồ có minh chứng tương ứng và mọi học sinh có cách tham gia an toàn, phù hợp.",
        "Tổ trưởng điều phối lịch và minh chứng": "Tổ trưởng điều phối lịch, mẫu bản đồ và công cụ minh chứng; giáo viên thiết kế nhiệm vụ, phân hóa và phản hồi; lãnh đạo phê duyệt an toàn; học sinh chọn vai trò, cập nhật bản đồ, thực hiện và phản tư; gia đình/cộng đồng phối hợp tự nguyện trong phạm vi được duyệt. Ưu tiên vật liệu sạch, công cụ sẵn có và phương án không dùng thiết bị.",
        "Áp phích đẹp chưa đủ chứng minh năng lực": "Bản đồ đẹp chưa đủ chứng minh năng lực. Hồ sơ cần cho thấy quá trình hình thành sản phẩm: thẻ “Một giọt nước kể chuyện”, phiếu dữ kiện – ý kiến, bảng nguyên nhân – hệ quả, bản nháp giải pháp, phân vai, nhật ký hành động, phản hồi và bài phản tư. Đánh giá theo bốn nhóm: hiểu biết và nguồn; đồng cảm và đối thoại; hành động an toàn, có căn cứ; hợp tác và trách nhiệm số.",
        "Kế hoạch đề xuất chuyển GCED từ một định hướng rộng": "Giá trị nổi bật của kế hoạch nằm ở việc biến GCED thành một sản phẩm học tập có thể nhìn thấy và tiếp tục phát triển. Qua “Bản đồ dòng chảy xanh”, học sinh không chỉ nói về bảo vệ nước mà phải kiểm tra thông tin, nhận diện nhiều góc nhìn, giải thích hệ quả, thử một hành động nhỏ và phản tư về điều cần sửa. Mỗi dấu vết trên bản đồ đồng thời là minh chứng để giáo viên phản hồi và điều chỉnh kế hoạch.",
        "Trước khi áp dụng, trường phải cập nhật bộ sách": "Trước khi áp dụng, trường phải cập nhật bộ sách, lịch năm học, đặc điểm người học, nhân lực, thiết bị, quy trình an toàn và đầu mối phối hợp. Bản đồ và thẻ học tập có thể thay đổi về hình thức, nhưng phải giữ các yêu cầu cốt lõi: bám YCCĐ, có tiếng nói của học sinh, có phương án tham gia thay thế, bảo vệ dữ liệu và chỉ kết luận hiệu quả từ minh chứng thực tế.",
    }
    for p in doc.paragraphs:
        stripped = p.text.strip()
        for start, new_text in replacements.items():
            if stripped.startswith(start):
                p.text = new_text
                break
        if stripped.startswith("Hình 1. Tiến trình triển khai kế hoạch giáo dục công dân toàn cầu"):
            label = "Hình 1. Tiến trình phát triển sản phẩm “Dòng chảy xanh” trong năm học"
            p.text = label if p.style.name == "Figure Caption" else label + "\t5"


doc = Document(SOURCE)
align_report_with_product(doc)
insert_product(doc)
update_front(doc, table1_page=10, table2_page=11, refs_page=13, figure2_page=9, figure3_page=9)
doc.save(OUTPUT)
print(OUTPUT)
